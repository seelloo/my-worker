"""Result file preview API — read Excel/text content for log result inspection."""
from flask import Blueprint, jsonify, request

from utils.security import validate_file_path

preview_bp = Blueprint("preview", __name__, url_prefix="/api/preview")


@preview_bp.route("/excel", methods=["GET"])
def api_preview_excel():
    path = request.args.get("path", "").strip()
    rows = request.args.get("rows", 20, type=int)

    is_valid, err_msg = validate_file_path(path, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    rows = max(1, min(rows, 100))

    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet = wb.active
        headers = []
        data_rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(c) if c is not None else "" for c in row]
            elif i <= rows:
                data_rows.append([str(c) if c is not None else "" for c in row])
            else:
                break
        wb.close()
        total_rows = sheet.max_row - 1 if sheet.max_row else 0
        return jsonify({
            "status": "success",
            "headers": headers,
            "rows": data_rows,
            "total_rows": total_rows,
            "sheet_name": sheet.title
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@preview_bp.route("/file", methods=["GET"])
def api_preview_file():
    path = request.args.get("path", "").strip()
    lines = request.args.get("lines", 50, type=int)

    is_valid, err_msg = validate_file_path(path, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    lines = max(1, min(lines, 200))

    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content_lines = [f.readline() for _ in range(lines)]
        content = "".join(content_lines)
        return jsonify({
            "status": "success",
            "content": content,
            "lines": len(content_lines)
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@preview_bp.route("/excel/sheets", methods=["GET"])
def api_excel_sheets():
    """返回指定 Excel 文件的所有 Sheet 名称列表。"""
    path = request.args.get("path", "").strip()
    if not path:
        return jsonify({"status": "error", "message": "缺少 path 参数"}), 400

    is_valid, err_msg = validate_file_path(path, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = wb.sheetnames
        wb.close()
        return jsonify({"status": "success", "sheets": sheets})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500
