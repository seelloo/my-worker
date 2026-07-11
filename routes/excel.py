import os
import sys
import json
import tempfile
from pathlib import Path
from flask import Blueprint, jsonify, request, send_file

from scripts.filter_excel_by_condition import load_conditions_from_excel
from scripts.cross_sheet_merge import get_sheet_names, get_sheet_columns
from scripts.excel_column_reorder import parse_txt_fields
from scripts.excel_xlookup import get_columns as xlookup_get_columns
from utils.flask_utils import stream_task
from utils.security import validate_file_path, validate_directory_path, require_params

excel_bp = Blueprint('excel', __name__, url_prefix='/api')


@excel_bp.route('/import_conditions', methods=['POST'])
def api_import_conditions():
    """从 Excel 条件文档导入筛选条件"""
    data = request.json or {}
    file_path = (data.get('path') or '').strip()

    # 安全验证
    is_valid, err_msg = validate_file_path(file_path, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    try:
        conditions = load_conditions_from_excel(file_path)
        return jsonify({"status": "success", "conditions": conditions})
    except (FileNotFoundError, ValueError) as e:
        return jsonify({"status": "error", "message": str(e)}), 422
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/export_conditions', methods=['POST'])
def api_export_conditions():
    """将筛选条件导出为 Excel 条件模板"""
    data = request.json or {}
    conditions = data.get('conditions', [])
    if not conditions:
        return jsonify({"status": "error", "message": "没有条件可导出"}), 400

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "筛选条件"
    ws.append(["目标列", "操作符", "比对值"])

    for cond in conditions:
        col = (cond.get("col") or "").strip()
        op_val = (cond.get("op") or "").strip()
        target_val = cond.get("val", "")
        ws.append([col, op_val, target_val if target_val is not None else ""])

    fd, temp_path = tempfile.mkstemp(suffix='.xlsx')
    os.close(fd)
    wb.save(temp_path)
    wb.close()

    return send_file(
        temp_path,
        as_attachment=True,
        download_name="filter_conditions.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@excel_bp.route('/excel/get_sheets', methods=['POST'])
def api_excel_get_sheets():
    """获取 Excel 文件的所有 Sheet 名称列表"""
    data = request.json or {}
    src_file = data.get('src_file', '').strip()

    # 安全验证
    is_valid, err_msg = validate_file_path(src_file, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    try:
        sheets = get_sheet_names(src_file)
        return jsonify({"status": "success", "sheets": sheets})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/get_sheet_columns', methods=['POST'])
def api_excel_get_sheet_columns():
    """获取指定 Sheet 的列信息（列字母 + 表头名称）"""
    data = request.json or {}
    src_file = data.get('src_file', '').strip()
    sheet_name = data.get('sheet_name', '').strip()

    # 安全验证
    is_valid, err_msg = validate_file_path(src_file, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    try:
        # sheet_name 为空时自动取第一个 Sheet
        if not sheet_name:
            sheet_name = get_sheet_names(src_file)[0]
        cols = get_sheet_columns(src_file, sheet_name)
        return jsonify({"status": "success", "columns": cols})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/cross_sheet_merge', methods=['POST'])
def api_excel_cross_sheet_merge():
    """执行跨 Sheet 多列合并（流式输出）"""
    data = request.json or {}
    src_file = data.get('src_file', '').strip()
    out_file = data.get('out_file', '').strip()
    out_sheet = data.get('out_sheet', '合并结果').strip()
    skip_header = data.get('skip_header', True)
    columns_cfg = data.get('columns', [])

    # 安全验证 - 验证源文件
    is_valid, err_msg = validate_file_path(src_file, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    if not out_file:
        return jsonify({"status": "error", "message": "缺少 out_file 参数"}), 400

    if not columns_cfg:
        return jsonify({"status": "error", "message": "columns 配置为空"}), 400

    fd, cfg_path = tempfile.mkstemp(suffix='.json')
    try:
        os.close(fd)
        with open(cfg_path, 'w', encoding='utf-8') as f:
            json.dump(columns_cfg, f, ensure_ascii=False)

        script = os.path.join(os.path.dirname(__file__), '..', 'scripts', 'cross_sheet_merge.py')
        cmd = [
            sys.executable, script,
            '-s', src_file,
            '-o', out_file,
            '-c', cfg_path,
            '--out_sheet', out_sheet,
        ]
        if not skip_header:
            cmd.append('--keep_header')

        return stream_task(cmd)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/split_sheets', methods=['POST'])
def api_excel_split_sheets():
    """将多 Sheet Excel 按指定列值拆分为多个文件（流式输出）"""
    data = request.json or {}
    input_path = data.get('input_path', '').strip()
    split_col = data.get('split_col', '').strip()
    output_dir = data.get('output_dir', '').strip()

    if not input_path or not split_col or not output_dir:
        return jsonify({"status": "error", "message": "参数缺失：input_path / split_col / output_dir"}), 400

    # 安全验证
    is_valid, err_msg = validate_file_path(input_path, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    is_valid, err_msg = validate_directory_path(output_dir, must_exist=False)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    script = os.path.join(os.path.dirname(__file__), '..', 'scripts', 'excel_sheet_splitter.py')
    cmd = [
        sys.executable, script,
        'split',
        '--input', input_path,
        '--col', split_col,
        '--output_dir', output_dir,
    ]
    return stream_task(cmd)


@excel_bp.route('/excel/merge_sheets', methods=['POST'])
def api_excel_merge_sheets():
    """将目录内同名 Excel 文件按 Sheet 合并重组（流式输出）"""
    data = request.json or {}
    input_dir = data.get('input_dir', '').strip()
    output_dir = data.get('output_dir', '').strip()

    if not input_dir or not output_dir:
        return jsonify({"status": "error", "message": "参数缺失：input_dir / output_dir"}), 400

    # 安全验证
    is_valid, err_msg = validate_directory_path(input_dir, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    is_valid, err_msg = validate_directory_path(output_dir, must_exist=False)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    script = os.path.join(os.path.dirname(__file__), '..', 'scripts', 'excel_sheet_splitter.py')
    cmd = [
        sys.executable, script,
        'merge',
        '--input_dir', input_dir,
        '--output_dir', output_dir,
    ]
    return stream_task(cmd)


@excel_bp.route('/excel/preview_split_values', methods=['POST'])
def api_excel_preview_split_values():
    """预览指定 Excel 文件某列的所有唯一值（轻量接口，非流式）"""
    data = request.json or {}
    input_path = data.get('input_path', '').strip()
    split_col = data.get('split_col', '').strip()

    if not input_path or not split_col:
        return jsonify({"status": "error", "message": "参数缺失"}), 400

    # 安全验证
    is_valid, err_msg = validate_file_path(input_path, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    try:
        import pandas as pd
        xl = pd.ExcelFile(input_path, engine='openpyxl')
        values = set()
        for sheet in xl.sheet_names:
            df = xl.parse(sheet, dtype=str)
            if split_col in df.columns:
                values.update(df[split_col].dropna().astype(str).unique())
        if not values:
            return jsonify({"status": "success", "values": [], "message": f"未在任何 Sheet 中找到列「{split_col}」"})
        return jsonify({"status": "success", "values": sorted(values)})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/group_assign', methods=['POST'])
def api_excel_group_assign():
    """按条件规则对数值列进行分组编码（流式输出）"""
    data = request.json or {}
    src_file = data.get('src_file', '').strip()
    out_file = data.get('out_file', '').strip()
    target_col = data.get('target_col', '').strip()
    group_col_name = data.get('group_col_name', '组编码').strip()
    concat_col = data.get('concat_col', '').strip()
    rules = data.get('rules', [])

    if not src_file or not out_file or not target_col:
        return jsonify({"status": "error", "message": "参数缺失：src_file / out_file / target_col"}), 400

    # 安全验证
    is_valid, err_msg = validate_file_path(src_file, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    if not rules:
        return jsonify({"status": "error", "message": "至少需要一条分组规则"}), 400

    fd, rules_path = tempfile.mkstemp(suffix='.json')
    try:
        os.close(fd)
        with open(rules_path, 'w', encoding='utf-8') as f:
            json.dump(rules, f, ensure_ascii=False)

        script = os.path.join(os.path.dirname(__file__), '..', 'scripts', 'excel_group_assign.py')
        cmd = [
            sys.executable, script,
            '-s', src_file,
            '-o', out_file,
            '-c', target_col,
            '-r', rules_path,
            '--group_col_name', group_col_name,
        ]
        if concat_col:
            cmd.extend(['--concat_col', concat_col])

        return stream_task(cmd)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/parse_reorder_txt', methods=['POST'])
def api_excel_parse_reorder_txt():
    """解析 TXT 文件，返回有序字段名列表（每行一个字段名）"""
    data = request.json or {}
    txt_path = (data.get('txt_path') or '').strip()

    if not txt_path:
        return jsonify({"status": "error", "message": "txt_path 参数不能为空"}), 400

    # TXT 文件不在通用扩展名白名单内，单独做路径安全 + 后缀校验
    from utils.security import validate_path
    is_valid, err_msg = validate_path(txt_path)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    if not txt_path.lower().endswith('.txt'):
        return jsonify({"status": "error", "message": "仅支持 .txt 格式的字段清单文件"}), 400

    import os
    if not os.path.isfile(txt_path):
        return jsonify({"status": "error", "message": f"文件不存在: {os.path.basename(txt_path)}"}), 400

    try:
        fields = parse_txt_fields(txt_path)
        return jsonify({"status": "success", "fields": fields, "count": len(fields)})
    except (FileNotFoundError, ValueError) as e:
        return jsonify({"status": "error", "message": str(e)}), 422
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/reorder_columns', methods=['POST'])
def api_excel_reorder_columns():
    """按指定字段顺序重组 Excel 列并导出到新文件（流式输出）"""
    data = request.json or {}
    src_file = data.get('src_file', '').strip()
    out_file = data.get('out_file', '').strip()
    sheet_name = data.get('sheet_name', '').strip()
    fields = data.get('fields', [])  # 有序字段名列表

    # 参数校验
    is_valid, err_msg = validate_file_path(src_file, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    if not out_file:
        return jsonify({"status": "error", "message": "缺少 out_file 参数"}), 400

    if not fields or not isinstance(fields, list):
        return jsonify({"status": "error", "message": "fields 参数为空或格式错误"}), 400

    fd, fields_path = tempfile.mkstemp(suffix='.json')
    try:
        os.close(fd)
        with open(fields_path, 'w', encoding='utf-8') as f:
            json.dump(fields, f, ensure_ascii=False)

        script = os.path.join(os.path.dirname(__file__), '..', 'scripts', 'excel_column_reorder.py')
        cmd = [
            sys.executable, script,
            '-s', src_file,
            '-o', out_file,
            '-f', json.dumps(fields, ensure_ascii=False),
        ]
        if sheet_name:
            cmd.extend(['--sheet', sheet_name])

        return stream_task(cmd)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/xlookup_get_cols', methods=['POST'])
def api_xlookup_get_cols():
    """获取指定 Excel 文件/Sheet 的列名列表，供前端渲染选项"""
    data = request.json or {}
    src_file   = (data.get('src_file')   or '').strip()
    sheet_name = (data.get('sheet_name') or '').strip()

    is_valid, err_msg = validate_file_path(src_file, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    try:
        if not sheet_name:
            sheet_name = get_sheet_names(src_file)[0]
        columns = xlookup_get_columns(src_file, sheet_name)
        return jsonify({"status": "success", "columns": columns})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/xlookup_run', methods=['POST'])
def api_xlookup_run():
    """执行 XLookup 合并（流式输出）"""
    data = request.json or {}
    main_file     = (data.get('main_file')     or '').strip()
    main_sheet    = (data.get('main_sheet')    or '').strip()
    lookup_col    = (data.get('lookup_col')    or '').strip()
    lookup_file   = (data.get('lookup_file')   or '').strip()
    lookup_sheet  = (data.get('lookup_sheet')  or '').strip()
    match_col     = (data.get('match_col')     or '').strip()
    result_cols   = data.get('result_cols', [])
    out_file      = (data.get('out_file')      or '').strip()
    count_col     = (data.get('count_col')     or '匹配记录数').strip()

    # 参数校验
    for label, path in [('主文件', main_file), ('副文件', lookup_file)]:
        is_valid, err_msg = validate_file_path(path, must_exist=True)
        if not is_valid:
            return jsonify({"status": "error", "message": f"{label}: {err_msg}"}), 400

    if not lookup_col:
        return jsonify({"status": "error", "message": "缺少 lookup_col 参数"}), 400
    if not match_col:
        return jsonify({"status": "error", "message": "缺少 match_col 参数"}), 400
    if not out_file:
        return jsonify({"status": "error", "message": "缺少 out_file 参数"}), 400
    if not result_cols or not isinstance(result_cols, list):
        return jsonify({"status": "error", "message": "result_cols 不能为空"}), 400

    try:
        script = os.path.join(os.path.dirname(__file__), '..', 'scripts', 'excel_xlookup.py')
        cmd = [
            sys.executable, script,
            '--main',         main_file,
            '--lookup_col',   lookup_col,
            '--lookup',       lookup_file,
            '--match_col',    match_col,
            '--result_cols',  json.dumps(result_cols, ensure_ascii=False),
            '--out',          out_file,
            '--count_col',    count_col,
        ]
        if main_sheet:
            cmd.extend(['--main_sheet', main_sheet])
        if lookup_sheet:
            cmd.extend(['--lookup_sheet', lookup_sheet])

        return stream_task(cmd)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/get_headers', methods=['POST'])
def api_excel_get_headers():
    """获取 Excel 指定 Sheet 的有序列名列表（用于字段选择器）"""
    data = request.json or {}
    src_file = (data.get('src_file') or '').strip()
    sheet_name = (data.get('sheet_name') or '').strip()

    is_valid, err_msg = validate_file_path(src_file, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    try:
        from openpyxl import load_workbook
        wb = load_workbook(src_file, data_only=True, read_only=True)
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active
            sheet_name = ws.title

        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        wb.close()

        if not header_row:
            return jsonify({"status": "error", "message": "表头为空，无法读取列名"}), 422

        headers = [str(h).strip() if h is not None else '' for h in header_row]
        named_headers = [h for h in headers if h]

        return jsonify({
            "status": "success",
            "headers": named_headers,
            "sheet": sheet_name,
            "total": len(named_headers),
        })
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@excel_bp.route('/excel/field_select_run', methods=['POST'])
def api_excel_field_select_run():
    """按选定字段顺序重组 Excel 列，自动推断输出路径，流式输出"""
    data = request.json or {}
    src_file   = (data.get('src_file')   or '').strip()
    sheet_name = (data.get('sheet_name') or '').strip()
    fields     = data.get('fields', [])

    is_valid, err_msg = validate_file_path(src_file, must_exist=True)
    if not is_valid:
        return jsonify({"status": "error", "message": err_msg}), 400

    if not fields or not isinstance(fields, list):
        return jsonify({"status": "error", "message": "fields 不能为空"}), 400

    # 自动推断输出路径：同目录，文件名加 _selected 后缀
    src_path = Path(src_file)
    out_file = str(src_path.parent / f"{src_path.stem}_selected{src_path.suffix}")

    try:
        import subprocess
        from flask import Response, stream_with_context

        script = os.path.join(os.path.dirname(__file__), '..', 'scripts', 'excel_column_reorder.py')
        cmd = [
            sys.executable, script,
            '-s', src_file,
            '-o', out_file,
            '-f', json.dumps(fields, ensure_ascii=False),
        ]
        if sheet_name:
            cmd.extend(['--sheet', sheet_name])

        def generate():
            yield f"[INFO] 输出文件：{out_file}\n"
            _env = os.environ.copy()
            _env['PYTHONIOENCODING'] = 'utf-8'
            _env['PYTHONUTF8'] = '1'
            proc = subprocess.Popen(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding='utf-8',
                errors='replace',
                env=_env,
            )
            for line in proc.stdout:
                yield line
            proc.wait()

        return Response(stream_with_context(generate()), mimetype='text/plain; charset=utf-8')
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500
