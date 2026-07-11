import pytest
import os
import openpyxl
from pathlib import Path
from utils.excel_utils import ExcelUtils
from scripts.rename_excel_sheets import rename_excel_sheets

def test_generate_new_sheet_name():
    existing = set()
    
    # 1. 正常去除 XXX
    name1 = ExcelUtils.generate_new_sheet_name("Report_XXX_2023", remove_str="XXX", add_str=None, add_position="back", existing_names=existing)
    assert name1 == "Report__2023"
    existing.add(name1)
    
    # 2. 正常前边增加 YYY
    name2 = ExcelUtils.generate_new_sheet_name("Sales", remove_str=None, add_str="YYY_", add_position="front", existing_names=existing)
    assert name2 == "YYY_Sales"
    existing.add(name2)
    
    # 3. 先去除 XXX，再增加 YYY 在后边 (优先去除，然后增加)
    name3 = ExcelUtils.generate_new_sheet_name("Test_XXX_Sheet", remove_str="XXX", add_str="_YYY", add_position="back", existing_names=existing)
    assert name3 == "Test__Sheet_YYY"
    existing.add(name3)
    
    # 4. 超长名字，触发从后往前截断（测试长度保畒28安全位）
    very_long_name = "这是一个超级无敌非常长长长长长长长长长长长长长长长超过了三十一个字的测试Sheet名称"
    name4 = ExcelUtils.generate_new_sheet_name(very_long_name, remove_str=None, add_str="追加", add_position="back", existing_names=existing)
    assert len(name4) <= 31
    assert name4 == (very_long_name + "追加")[:28]  # 会被截断到28以备未来加后缀
    existing.add(name4)
    
    # 5. 引发截断后的名字冲突，触发后缀加成并依然不超过 31
    name5 = ExcelUtils.generate_new_sheet_name(very_long_name, remove_str=None, add_str="追加", add_position="back", existing_names=existing)
    assert len(name5) <= 31
    assert name5.endswith("_1")
    assert name5 != name4
    existing.add(name5)

def test_rename_excel_sheets(tmp_path, capsys):
    src_dir = tmp_path / "src"
    dest_dir = tmp_path / "dest"
    src_dir.mkdir()
    
    # 准备测试文件
    file_path = src_dir / "test_book.xlsx"
    wb = openpyxl.Workbook()
    
    # 创建4个不同情况的 Sheet 
    wb.active.title = "Test_XXX_1"   # 去除场景
    wb.create_sheet(title="Test_2")  # 追加场景
    wb.create_sheet(title="Data_XXX_3") # 综合场景
    
    # 截断和冲突场景 (让前 28 位之后才出现差异，从而在截断时丢失差异引发重名)
    long_name = "这是一个专门为了测试前端加长导致截断和冲突的极长极长名字_XXX"
    wb.create_sheet(title=long_name + "结尾差异A")
    wb.create_sheet(title=long_name + "结尾差异B")
    
    wb.save(file_path)
    
    # 执行重命名：去除 XXX，前缀加 NEW_
    # 此时 src 参数直接传入文件路径 file_path 即可，目标依然填 dest_dir，测试其推断逻辑
    rename_excel_sheets(str(file_path), str(dest_dir), remove_str="XXX", add_str="NEW_", add_position="front")
    
    # 验证有没有顺利另存（而不覆盖）
    out_file = dest_dir / "test_book.xlsx"
    assert out_file.exists()
    
    # 读取输出的文件验证名字
    wb_out = openpyxl.load_workbook(out_file)
    names = wb_out.sheetnames
    
    # "Test_XXX_1" -> 去除XXX -> "Test__1" -> 前面加 NEW_ -> "NEW_Test__1"
    assert "NEW_Test__1" in names
    
    # "Test_2" -> 去除XXX无效 -> 前面加 NEW_ -> "NEW_Test_2"
    assert "NEW_Test_2" in names
    
    # "Data_XXX_3" -> "NEW_Data__3"
    assert "NEW_Data__3" in names
    
    # 极长名字测试，因为前缀 NEW_加进去绝对超长
    # 两个带有 XXX 的会在去除了 XXX 之后，加上前缀 NEW_ 然后被截断并分配出后缀 "_1"
    # 我们预期至少有一个带后缀并符合总长度限制
    long_names_found = [n for n in names if n.startswith("NEW_这是一个专门为了测试前端加长导")]
    assert len(long_names_found) == 2
    assert all(len(n) <= 31 for n in long_names_found)
    assert any(n.endswith("_1") for n in long_names_found)
    
    wb_out.close()
    
    # 验证有没有打印对应的提示语句
    captured = capsys.readouterr()
    assert "因新名字" in captured.out
    assert "超过31个字，已从后往前截断" in captured.out
