import os
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import csv
import pytest
from openpyxl import load_workbook
from scripts.csv_to_excel_text import convert_csvs_to_excel

@pytest.fixture
def dummy_csv_dir(tmpdir):
    # Dummy valid csv 1 (UTF-8)
    file1 = tmpdir.join("test1.csv")
    with open(str(file1), 'w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(["ID", "Name", "Number"])
        writer.writerow(["001", "Alice", "1E5"]) # should not become 100000 or exponent in excel
    
    # Dummy valid csv 2 (GBK encoded)
    file2 = tmpdir.join("test2.csv")
    with open(str(file2), 'w', encoding='gbk', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(["编号", "产品"])
        writer.writerow(["000555", "苹果"]) # leading zeros should be kept
        
    return str(tmpdir)

def test_convert_csvs_to_excel(dummy_csv_dir, tmpdir):
    out_file = os.path.join(str(tmpdir), "output.xlsx")
    
    success = convert_csvs_to_excel(dummy_csv_dir, out_file)
    assert success is True
    assert os.path.exists(out_file)
    
    wb = load_workbook(out_file)
    # Check sheets exist with correct names
    assert "test1" in wb.sheetnames
    assert "test2" in wb.sheetnames
    
    # Validate test1 (UTF-8)
    ws1 = wb["test1"]
    assert ws1["A2"].value == "001"
    assert ws1["A2"].number_format == '@'
    assert ws1["C2"].value == "1E5"
    assert ws1["C2"].number_format == '@'
    
    # Validate test2 (GBK)
    ws2 = wb["test2"]
    assert ws2["A2"].value == "000555"
    assert ws2["A2"].number_format == '@'
    assert ws2["B2"].value == "苹果"
