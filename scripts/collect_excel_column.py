from utils.file_utils import FileUtils
from utils.excel_utils import ExcelUtils
from utils.error_handler import ErrorHandler
from utils.cli_parser import CLIParser
import argparse
import sys
from pathlib import Path
import openpyxl

def collect_column_data(src_file: str, dest_path_str: str, target_column: str, out_sheet_name: str = "Collected_Data"):
    src_path = FileUtils.resolve_path(src_file)
    dest_path = FileUtils.resolve_path(dest_path_str)

    if not FileUtils.validate_file_exists(src_path):
        ErrorHandler.handle_file_not_found(src_file)
        return

    out_file_path = FileUtils.get_output_path(src_path, dest_path)
    out_file_path = FileUtils.handle_existing_file(out_file_path)

    try:
        wb = openpyxl.load_workbook(src_path)
    except Exception as e:
        ErrorHandler.handle_file_read_error(src_path.name, e)
        return

    ErrorHandler.print_info(f"正在处理: '{src_path.name}'，目标列: {target_column} 列")

    # 集合用于自动去重
    collected_data = set()
    total_cells_scanned = 0
    sheets_scanned = 0

    # 获取所有的 Sheet 名称，避免在遍历中动态添加导致死循环
    original_sheet_names = wb.sheetnames

    # 如果源文件里已经有我们想要创建名字的 Sheet，说明可能是重复执行，或者是原表带有的
    # 我们应该跳过它不提取数据（如果它是我们要输出的表），当然名字起冲突后面会处理

    for sheet_name in original_sheet_names:
        # 如果原始表格中已经有同名的收集表，我们不要提取它自己的内容
        if sheet_name == out_sheet_name:
            continue

        ws = wb[sheet_name]
        sheets_scanned += 1

        # 从单个Sheet中收集指定列的数据
        sheet_data = ExcelUtils.collect_column_data_from_sheet(ws, target_column)
        collected_data.update(sheet_data)
        total_cells_scanned += len(sheet_data)

    ErrorHandler.print_success(f"扫描了 {sheets_scanned} 个 Sheet，共 {total_cells_scanned} 个数据单元格(跳过表头)。")
    ErrorHandler.print_success(f"经过分析去重后，提炼出 {len(collected_data)} 条唯一的非空数据。")

    # 建立我们汇总用的新 Sheet
    # 名字处理逻辑（防冲突）：
    final_out_sheet_name = ExcelUtils.get_valid_sheet_name(out_sheet_name, set(wb.sheetnames))

    out_ws = ExcelUtils.create_output_sheet(wb, final_out_sheet_name)

    # 加上我们这唯一一列的表头（可选，但通常汇总完最好有个头，既然源的也跳过了头）
    out_ws.cell(row=1, column=1, value="汇总去重数据")

    # 把集合内容转为列表（为了稍微排个序，可选项。字符串字典序排一下好看）
    sorted_data = sorted(list(collected_data))

    # 从第二行开始，全部写到 A列 (column 1)
    current_row = 2
    for item in sorted_data:
        out_ws.cell(row=current_row, column=1, value=item)
        current_row += 1

    ErrorHandler.print_success(f"汇总数据成功写入了名为 '{final_out_sheet_name}' 的新 Sheet 的 A 列。")

    try:
        ExcelUtils.save_workbook(wb, out_file_path)
        ErrorHandler.print_success(f"🎉 处理完成！文件已另存为 -> '{out_file_path.name}'\n")
    except Exception as e:
        ErrorHandler.handle_file_save_error(out_file_path.name, e)

def main():
    parser = CLIParser.create_parser(
        description="收集 Excel (.xlsx) 指定单文件中所有 Sheet 的某一列内容，强制跳过表头，去重后汇集到一个新的 Sheet 中的 A 列。"
    )
    CLIParser.add_common_arguments(parser)
    CLIParser.add_column_arguments(parser)
    args = CLIParser.parse_arguments(parser)

    CLIParser.validate_column_argument(args.column)

    # 规范化列名为大写
    collect_column_data(args.file, args.out, args.column.upper(), args.sheet)

if __name__ == "__main__":
    main()
