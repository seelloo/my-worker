"""
pytest 测试：accounts_daily_report.py
======================================
覆盖场景：
  1. is_nested_list() — 嵌套/非嵌套判断
  2. copy_excel_file() — openpyxl 拷贝文件
  3. accounts_daily_report() — 正常流程（新建文件）
  4. accounts_daily_report() — 文件已存在（追加写入）
  5. main() CLI — 正常流程端到端验证
  6. main() CLI — 缺失列名的错误提示
"""

import sys
import os
import pytest
import pandas as pd
import openpyxl

# 确保从项目根目录可以引入 scripts 模块
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from scripts.accounts_daily_report import (
    is_nested_list,
    copy_excel_file,
    accounts_daily_report,
    main,
)


# ---------------------------------------------------------------------------
# 工具函数测试
# ---------------------------------------------------------------------------

class TestIsNestedList:
    def test_nested_returns_true(self):
        assert is_nested_list([[1, 2], [3, 4]]) is True

    def test_single_nested_returns_true(self):
        assert is_nested_list(["a", ["b", "c"]]) is True

    def test_flat_returns_false(self):
        assert is_nested_list(["a", "b", "c"]) is False

    def test_empty_returns_false(self):
        assert is_nested_list([]) is False

    def test_numbers_returns_false(self):
        assert is_nested_list([1, 2, 3]) is False


class TestCopyExcelFile:
    def test_copy_creates_target(self, tmp_path):
        # 创建源文件
        src = tmp_path / "source.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "测试数据"
        wb.save(str(src))

        dst = tmp_path / "target.xlsx"
        copy_excel_file(str(src), str(dst))

        assert dst.exists()
        wb2 = openpyxl.load_workbook(str(dst))
        assert wb2.active["A1"].value == "测试数据"


# ---------------------------------------------------------------------------
# 核心业务函数测试
# ---------------------------------------------------------------------------

def _build_daily_excel(path, skiprows=1):
    """构造佣金支付日报表测试数据（含跳过行）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    # 第1行是脏数据（skiprows=1 时会被跳过）
    ws.append(["脏数据行"])
    # 第2行是表头
    ws.append(["财辅报账单号", "实付金额(元)", "供应商名称"])
    # 数据行
    ws.append(["BZ001", 1000, "供应商A"])
    ws.append(["BZ002", 2000, "供应商B"])
    ws.append(["BZ003", 3000, "供应商C"])
    wb.save(str(path))


def _build_template_excel(path):
    """构造电子签模板文件（Sheet1 前3行为模板头）。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "电子签报表标题"
    ws["A2"] = "公司信息"
    ws["A3"] = "表头行"
    wb.save(str(path))


class TestAccountsDailyReport:
    def test_new_file_created_from_template(self, tmp_path):
        """正常流程：输出文件不存在时，自动拷贝模板并写入数据。"""
        daily_file = tmp_path / "佣金支付日报表.xlsx"
        template_file = tmp_path / "模板.xlsx"
        _build_daily_excel(daily_file)
        _build_template_excel(template_file)

        result = accounts_daily_report(
            bzbm_list=["BZ001", "BZ002"],
            base_path=str(tmp_path),
            accounts_daily_file="佣金支付日报表.xlsx",
            sign_report_file="模板.xlsx",
            id_col="财辅报账单号",
            amount_col="实付金额(元)",
        )

        # 总金额 = 1000 + 2000 = 3000
        assert result["total_amount"] == 3000
        assert result["matched_rows"] == 2
        assert result["file_existed"] is False
        # 输出文件名包含总金额
        assert "3000" in os.path.basename(result["output_file"])
        assert os.path.exists(result["output_file"])

    def test_existing_file_appends_data(self, tmp_path):
        """文件已存在时，直接 overlay 写入，is_exist 应为 True。"""
        daily_file = tmp_path / "佣金支付日报表.xlsx"
        template_file = tmp_path / "模板.xlsx"
        _build_daily_excel(daily_file)
        _build_template_excel(template_file)

        # 第一次运行，新建文件
        result1 = accounts_daily_report(
            bzbm_list=["BZ003"],
            base_path=str(tmp_path),
            accounts_daily_file="佣金支付日报表.xlsx",
            sign_report_file="模板.xlsx",
            id_col="财辅报账单号",
            amount_col="实付金额(元)",
        )
        assert result1["file_existed"] is False

        # 第二次运行（相同单号列表 → 相同总金额 → 相同文件名）
        result2 = accounts_daily_report(
            bzbm_list=["BZ003"],
            base_path=str(tmp_path),
            accounts_daily_file="佣金支付日报表.xlsx",
            sign_report_file="模板.xlsx",
            id_col="财辅报账单号",
            amount_col="实付金额(元)",
        )
        assert result2["file_existed"] is True
        assert result2["total_amount"] == 3000

    def test_unmatched_ids_return_zero(self, tmp_path):
        """单号无匹配时，总金额应为 0，匹配行数为 0。"""
        daily_file = tmp_path / "佣金支付日报表.xlsx"
        template_file = tmp_path / "模板.xlsx"
        _build_daily_excel(daily_file)
        _build_template_excel(template_file)

        result = accounts_daily_report(
            bzbm_list=["BZ999"],  # 不存在的单号
            base_path=str(tmp_path),
            accounts_daily_file="佣金支付日报表.xlsx",
            sign_report_file="模板.xlsx",
            id_col="财辅报账单号",
            amount_col="实付金额(元)",
        )

        assert result["total_amount"] == 0
        assert result["matched_rows"] == 0


# ---------------------------------------------------------------------------
# CLI 端到端测试
# ---------------------------------------------------------------------------

def _build_bzlist_excel(path, group_col="是否已报账", id_col="财辅报账单号"):
    """构造报账单台账测试数据。"""
    df = pd.DataFrame({
        group_col: ["已报账", "已报账", "未报账"],
        id_col: ["BZ001", "BZ002", "BZ003"],
    })
    df.to_excel(str(path), index=False)


class TestMainCLI:
    def test_full_flow(self, tmp_path, monkeypatch, capsys):
        """端到端：正确参数 → 正常执行并打印结果。"""
        bzlist = tmp_path / "bzlist.xlsx"
        daily = tmp_path / "佣金支付日报表.xlsx"
        template = tmp_path / "模板.xlsx"
        _build_bzlist_excel(bzlist)
        _build_daily_excel(daily)
        _build_template_excel(template)

        monkeypatch.setattr(
            sys, "argv",
            [
                "accounts_daily_report.py",
                "-b", str(bzlist),
                "-s", str(tmp_path),
                "-d", "佣金支付日报表.xlsx",
                "-t", "模板.xlsx",
            ],
        )
        main()
        captured = capsys.readouterr()
        assert "✅ 全部分组处理完毕" in captured.out

    def test_missing_group_col_exits_gracefully(self, tmp_path, monkeypatch, capsys):
        """错误列名 → 打印友好提示，不抛异常。"""
        bzlist = tmp_path / "bzlist.xlsx"
        daily = tmp_path / "佣金支付日报表.xlsx"
        template = tmp_path / "模板.xlsx"
        _build_bzlist_excel(bzlist)
        _build_daily_excel(daily)
        _build_template_excel(template)

        monkeypatch.setattr(
            sys, "argv",
            [
                "accounts_daily_report.py",
                "-b", str(bzlist),
                "-s", str(tmp_path),
                "-d", "佣金支付日报表.xlsx",
                "-t", "模板.xlsx",
                "-g", "不存在的列",  # 错误列名
            ],
        )
        main()
        captured = capsys.readouterr()
        assert "不存在分组列" in captured.out

    def test_nonexistent_bzlist_exits_gracefully(self, tmp_path, monkeypatch, capsys):
        """台账文件不存在 → 打印友好提示，不抛异常。"""
        monkeypatch.setattr(
            sys, "argv",
            [
                "accounts_daily_report.py",
                "-b", str(tmp_path / "不存在.xlsx"),
                "-s", str(tmp_path),
                "-d", "日报.xlsx",
                "-t", "模板.xlsx",
            ],
        )
        main()
        captured = capsys.readouterr()
        assert "不存在" in captured.out
