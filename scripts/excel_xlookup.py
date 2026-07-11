"""
excel_xlookup.py
-----------------
从副文件按指定列匹配主文件的查找列，将副文件的查询列（result_cols）
和一列「匹配记录数」追加到主文件末尾，输出新 Excel 文件。

- 匹配策略：转字符串 + 去首尾空格，大小写敏感，取第一条匹配
- 未匹配：result_cols 全填空，匹配记录数为 0
- 所有输出单元格强制文本格式 ('@')
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook, Workbook
from openpyxl.styles import numbers as xl_numbers


# ────────────────────────────────────────────────────────────────
# 辅助：获取列名列表
# ────────────────────────────────────────────────────────────────

def get_columns(src_file: str, sheet_name: str = None) -> list[str]:
    """
    返回指定文件第一行的列名列表（字符串）。
    sheet_name 为空时取第一个 Sheet。
    """
    wb = load_workbook(src_file, read_only=True, data_only=True)
    ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active
    header = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in next(ws.iter_rows(min_row=1, max_row=1), [])
    ]
    wb.close()
    return header


# ────────────────────────────────────────────────────────────────
# 核心：构建查找表
# ────────────────────────────────────────────────────────────────

def _build_lookup_table(
    lookup_file: str,
    lookup_sheet: str,
    match_col: str,
    result_cols: list[str],
) -> dict[str, dict]:
    """
    读取副文件，构建查找字典。
    key   = match_col 的值（str，strip）
    value = {
        "first": {col: value, ...},   # 第一条匹配行的 result_cols 值
        "count": int,                  # 匹配记录数
    }
    """
    wb = load_workbook(lookup_file, read_only=True, data_only=True)
    ws = wb[lookup_sheet] if (lookup_sheet and lookup_sheet in wb.sheetnames) else wb.active

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
    header = [str(h).strip() if h is not None else "" for h in header_row]

    # 列索引定位
    try:
        match_idx = header.index(match_col)
    except ValueError:
        wb.close()
        raise ValueError(f"副文件中找不到对应列「{match_col}」，请检查列名是否正确。")

    result_idxs = {}
    missing = []
    for col in result_cols:
        try:
            result_idxs[col] = header.index(col)
        except ValueError:
            missing.append(col)
    if missing:
        print(f"[WARN] 副文件中以下查询列不存在，将以空列代替: {missing}")

    # 统计 count（所有出现次数）
    count_map: dict[str, int] = defaultdict(int)
    first_map: dict[str, dict] = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        key_raw = row[match_idx] if match_idx < len(row) else None
        key = str(key_raw).strip() if key_raw is not None else ""
        count_map[key] += 1
        if key not in first_map:
            first_map[key] = {
                col: (str(row[idx]).strip() if idx < len(row) and row[idx] is not None else "")
                for col, idx in result_idxs.items()
            }
            # 不存在的列填空
            for col in missing:
                first_map[key][col] = ""

    wb.close()

    # 合并成最终查找表
    lookup: dict[str, dict] = {}
    for key in set(list(count_map.keys()) + list(first_map.keys())):
        lookup[key] = {
            "first": first_map.get(key, {col: "" for col in result_cols}),
            "count": count_map.get(key, 0),
        }

    print(f"[INFO] 副文件构建查找表完成，共 {len(lookup)} 个唯一 key。")
    return lookup


# ────────────────────────────────────────────────────────────────
# 核心：执行 XLookup 合并
# ────────────────────────────────────────────────────────────────

def xlookup_merge(
    main_file: str,
    main_sheet: str,
    lookup_col: str,
    lookup_file: str,
    lookup_sheet: str,
    match_col: str,
    result_cols: list[str],
    out_file: str,
    match_count_col_name: str = "匹配记录数",
):
    """
    将副文件的 result_cols（按 match_col 匹配主文件 lookup_col）追加到主文件，
    同时追加一列「匹配记录数」，输出新 Excel 文件。
    所有输出单元格强制文本格式。
    """
    main_path = Path(main_file).resolve()
    out_path = Path(out_file).resolve()

    if not main_path.exists():
        print(f"[ERROR] 主文件不存在: {main_path}")
        sys.exit(1)
    if not Path(lookup_file).resolve().exists():
        print(f"[ERROR] 副文件不存在: {lookup_file}")
        sys.exit(1)
    if not result_cols:
        print("[ERROR] 查询列（result_cols）不能为空。")
        sys.exit(1)

    # 构建查找表
    print(f"[INFO] 正在读取副文件构建查找表...")
    lookup_table = _build_lookup_table(lookup_file, lookup_sheet, match_col, result_cols)

    # 读取主文件
    print(f"[INFO] 正在读取主文件: {main_path.name}")
    wb_main = load_workbook(main_path, read_only=True, data_only=True)
    ws_main = wb_main[main_sheet] if (main_sheet and main_sheet in wb_main.sheetnames) else wb_main.active
    print(f"[INFO] 主文件 Sheet: {ws_main.title}")

    main_header_row = next(ws_main.iter_rows(min_row=1, max_row=1, values_only=True), [])
    main_header = [str(h).strip() if h is not None else "" for h in main_header_row]

    try:
        lookup_col_idx = main_header.index(lookup_col)
    except ValueError:
        wb_main.close()
        print(f"[ERROR] 主文件中找不到查找列「{lookup_col}」")
        sys.exit(1)

    # 创建输出文件
    wb_out = Workbook(write_only=True)
    ws_out = wb_out.create_sheet("XLookup结果")
    TEXT_FMT = xl_numbers.FORMAT_TEXT  # '@'

    def text_cell(value):
        from openpyxl.cell import WriteOnlyCell
        cell = WriteOnlyCell(ws_out, value=str(value) if value is not None else "")
        cell.number_format = TEXT_FMT
        return cell

    # 写入表头：主文件所有列 + result_cols + 匹配记录数
    out_header = list(main_header) + list(result_cols) + [match_count_col_name]
    ws_out.append([text_cell(h) for h in out_header])

    # 逐行处理
    row_count = 0
    matched_count = 0
    for row_values in ws_main.iter_rows(min_row=2, values_only=True):
        row_count += 1

        # 主文件原始列
        out_row = [text_cell(v) for v in row_values]
        # 补齐（若某行列数少于 header）
        while len(out_row) < len(main_header):
            out_row.append(text_cell(""))

        # 查找
        key_raw = row_values[lookup_col_idx] if lookup_col_idx < len(row_values) else None
        key = str(key_raw).strip() if key_raw is not None else ""
        match = lookup_table.get(key)

        if match:
            matched_count += 1
            for col in result_cols:
                out_row.append(text_cell(match["first"].get(col, "")))
            out_row.append(text_cell(str(match["count"])))
        else:
            for _ in result_cols:
                out_row.append(text_cell(""))
            out_row.append(text_cell("0"))

        ws_out.append(out_row)

        if row_count % 5000 == 0:
            print(f"[INFO] 已处理 {row_count} 行...")

    wb_main.close()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    print(f"[INFO] 正在保存至: {out_path.name} ...")
    wb_out.save(out_path)

    print(
        f"[DONE] 处理完成！共 {row_count} 行，其中 {matched_count} 行匹配成功，"
        f"{row_count - matched_count} 行未匹配。文件已保存至: {out_path}"
    )


# ────────────────────────────────────────────────────────────────
# 命令行入口
# ────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Excel XLookup 列匹配合并工具")
    parser.add_argument("--main",         required=True, help="主文件路径")
    parser.add_argument("--main_sheet",   default="",    help="主文件 Sheet 名（留空取第一个）")
    parser.add_argument("--lookup_col",   required=True, help="主文件查找列名")
    parser.add_argument("--lookup",       required=True, help="副文件路径")
    parser.add_argument("--lookup_sheet", default="",    help="副文件 Sheet 名（留空取第一个）")
    parser.add_argument("--match_col",    required=True, help="副文件对应列名")
    parser.add_argument("--result_cols",  required=True, help="副文件查询列名 JSON 数组，如 '[\"姓名\",\"部门\"]'")
    parser.add_argument("--out",          required=True, help="输出文件路径")
    parser.add_argument("--count_col",    default="匹配记录数", help="匹配记录数列名（默认：匹配记录数）")
    args = parser.parse_args()

    try:
        result_cols = json.loads(args.result_cols)
    except json.JSONDecodeError as e:
        print(f"[ERROR] result_cols JSON 解析失败: {e}")
        sys.exit(1)

    xlookup_merge(
        main_file=args.main,
        main_sheet=args.main_sheet,
        lookup_col=args.lookup_col,
        lookup_file=args.lookup,
        lookup_sheet=args.lookup_sheet,
        match_col=args.match_col,
        result_cols=result_cols,
        out_file=args.out,
        match_count_col_name=args.count_col,
    )
