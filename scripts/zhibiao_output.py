import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import warnings
warnings.filterwarnings('ignore')

import argparse

def main():
    parser = argparse.ArgumentParser(description="指标输出处理")
    parser.add_argument("-i", "--input", required=True, help="输入文件路径")
    parser.add_argument("-o", "--output", required=True, help="输出文件路径")
    parser.add_argument("-f", "--format", default="默认长表", help="输出格式：默认长表/套用模板宽表")
    parser.add_argument("-t", "--template", help="宽表模板文件路径")
    args = parser.parse_args()

    INPUT_FILE = args.input
    OUTPUT_FILE = args.output

    # ── 读取厅店表 ──────────────────────────────────────────────────
    df_store = pd.read_excel(INPUT_FILE, sheet_name='厅店', dtype=str)
    df_store.columns = df_store.columns.str.strip()
    
    # ── 读取指标表（双行表头：row0=大类, row1=子类）────────────────
    raw = pd.read_excel(INPUT_FILE, sheet_name='指标', header=None, dtype=str)
    
    # row0: 大类（月度目标值/权重/封顶值/保底值）
    # row1: 指标名称
    # row0~1 第0列 = "类型"
    big_cats = raw.iloc[0].ffill().tolist()   # 前向填充大类
    sub_cats  = raw.iloc[1].tolist()                          # 指标名
    
    # 构造多级列：(大类, 指标名)
    multi_cols = list(zip(big_cats, sub_cats))
    
    # 数据行从第2行开始，第0列是类型标签
    df_ind = raw.iloc[2:].reset_index(drop=True)
    df_ind.columns = pd.MultiIndex.from_tuples(multi_cols)
    
    # 第0列是类型名（用 ("类型","类型") 或第一个多级key）
    type_col_key = multi_cols[0]   # ('类型', '类型')
    df_ind.index = df_ind[type_col_key].str.strip().values
    df_ind = df_ind.drop(columns=[type_col_key])
    
    # 指标名列表（去掉第0列后，从 sub_cats[1:]）
    indicator_names = [s for s in sub_cats[1:] if pd.notna(s)]
    # 去重保序
    seen = set()
    unique_indicators = []
    for name in indicator_names:
        if name not in seen:
            unique_indicators.append(name)
            seen.add(name)

    if getattr(args, 'format', '默认长表') == '套用模板宽表' and getattr(args, 'template', None):
        import shutil
        import os
        if not os.path.exists(args.template):
            raise FileNotFoundError(f"未找到模板文件: {args.template}")
            
        shutil.copy(args.template, OUTPUT_FILE)
        wb = load_workbook(OUTPUT_FILE)
        ws = wb.active
        
        # 解析模板表头
        col_map = {} 
        current_l1 = None
        for c in range(1, ws.max_column + 1):
            v1 = ws.cell(row=1, column=c).value
            v2 = ws.cell(row=2, column=c).value
            if v1 is not None and str(v1).strip():
                current_l1 = str(v1).strip()
            l2 = str(v2).strip() if v2 is not None else ""
            if current_l1:
                col_map[(current_l1, l2)] = c
                
        # 清除原有数据 (第3行及以后)
        if ws.max_row > 2:
            ws.delete_rows(3, ws.max_row - 2)
            
        row_idx = 3
        # ── 整合数据并填入宽表 ──────────────────────────────────────────────────
        for _, store_row in df_store.iterrows():
            store_type = str(store_row.get('类型组合', '')).strip()
            actual_type = None
            if store_type in df_ind.index:
                actual_type = store_type
            else:
                for ind_k in df_ind.index:
                    if store_type.endswith(ind_k) and store_type == ind_k[:len(store_type)-len(ind_k)] + ind_k:
                        actual_type = ind_k
                        break
            if not actual_type:
                continue
                
            ind_row = df_ind.loc[actual_type]

            def get_val(big, ind):
                try:
                    v = ind_row[(big, ind)]
                    return float(v) if pd.notna(v) else 0
                except Exception:
                    return 0

            for (l1, l2), c in col_map.items():
                val = None
                if l2 == "" or l2 == l1 or str(l2).startswith("Unnamed"):
                    if l1 in store_row.index:
                        val = store_row[l1]
                else:
                    if l1 in unique_indicators:
                        if l2 == '目标值':
                            val = get_val('月度目标值', l1)
                        elif l2 == '权重':
                            val = get_val('权重', l1)
                        elif l2 in ['完成值', '完成率', '得分']:
                            val = 0
                            
                if l1 == '合计（保底60分）':
                    val = 60
                elif l1 == '总分' and val is None:
                    val = ''
                    
                if val is not None:
                    try:
                        if isinstance(val, str) and not val.strip():
                            pass
                        elif l2 in ['目标值', '权重', '完成值', '完成率', '得分']:
                            val = float(val) if '.' in str(val) else int(val)
                    except:
                        pass
                    ws.cell(row=row_idx, column=c).value = val
                    
            row_idx += 1

        wb.save(OUTPUT_FILE)
        print(f"完成！套用模板宽表，共输出 {row_idx - 3} 行数据 → {OUTPUT_FILE}")
        return

    
    # ── 整合 ─────────────────────────────────────────────────────────
    records = []
    for _, store_row in df_store.iterrows():
        store_type = str(store_row['类型组合']).strip()
        
        actual_type = None
        if store_type in df_ind.index:
            actual_type = store_type
        else:
            # 兼容处理“类型”被重复拼接的情况（如 终端厅终端厅城市支局 -> 终端厅城市支局）
            for ind_k in df_ind.index:
                if store_type.endswith(ind_k) and store_type == ind_k[:len(store_type)-len(ind_k)] + ind_k:
                    actual_type = ind_k
                    break
                    
        if not actual_type:
            continue
            
        ind_row = df_ind.loc[actual_type]
    
        for ind_name in unique_indicators:
            def get_val(big, ind):
                try:
                    v = ind_row[(big, ind)]
                    return v if pd.notna(v) else 0
                except Exception:
                    return 0
    
            records.append({
                '账期月份':     store_row['帐期'],
                '地市':         store_row['地市'],
                '店中商编码':   store_row['店中商编码'],
                '店中商名称':   store_row['店中商名称'],
                '指标名称':     ind_name,
                '月度目标值':   get_val('月度目标值', ind_name),
                '权重':         get_val('权重',    ind_name),
                '封顶值':       get_val('封顶值',  ind_name),
                '保底值':       get_val('保底值',  ind_name),
            })
    
    df_out = pd.DataFrame(records)
    
    # 数值列转换
    for col in ['月度目标值', '权重', '封顶值', '保底值']:
        df_out[col] = pd.to_numeric(df_out[col], errors='coerce').fillna(0)
    
    # ── 写出 Excel ────────────────────────────────────────────────────
    df_out.to_excel(OUTPUT_FILE, index=False, sheet_name='输出')
    
    # ── 美化格式 ─────────────────────────────────────────────────────
    wb = load_workbook(OUTPUT_FILE)
    ws = wb.active
    
    header_fill = PatternFill('solid', start_color='366092', end_color='366092')
    header_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    data_font   = Font(name='Arial', size=10)
    center      = Alignment(horizontal='center', vertical='center')
    left        = Alignment(horizontal='left',   vertical='center')
    thin        = Side(style='thin', color='CCCCCC')
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)
    
    col_widths = [12, 8, 18, 22, 22, 12, 8, 10, 8]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    
    for cell in ws[1]:
        cell.fill      = header_fill
        cell.font      = header_font
        cell.alignment = center
        cell.border    = border
    
    num_cols = {6, 7, 8, 9}   # 月度目标值/权重/封顶值/保底值
    alt_fill = PatternFill('solid', start_color='EBF3F9', end_color='EBF3F9')
    
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, cell in enumerate(row, 1):
            cell.font   = data_font
            cell.border = border
            if fill:
                cell.fill = fill
            if col_idx in num_cols:
                cell.alignment = center
                cell.number_format = 'General'
            else:
                cell.alignment = left
    
    ws.freeze_panes = 'A2'
    wb.save(OUTPUT_FILE)

    print(f"完成！共输出 {len(df_out)} 行数据 → {OUTPUT_FILE}")
    print(f"涉及店铺：{df_out['店中商编码'].nunique()} 家，指标数：{df_out['指标名称'].nunique()} 个")

if __name__ == '__main__':
    main()