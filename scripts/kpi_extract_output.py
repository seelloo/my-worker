"""
kpi_extract_output.py
从"指标" sheet（宽表）提取门店KPI指标数据，
输出为"输出" sheet 格式的长表 Excel。
账期月份 和 地市 两个字段由命令行参数手工传入。

Sheet 结构（基于示例文件）：
  - Sheet0="指标"（宽表）:
      行0：大类（月度指标值/权重/封顶值），前11列为 NaN（门店基础信息）
      行1：列名（区域/经营主体编码/.../店中商编码/店中商名称/... 及各指标名）
      行2+：数据行
  - Sheet1="输出"（长表）:
      列：账期月份, 地市, 店中商编码, 店中商名称, 指标名称, 月度目标值, 权重, 封顶值, 保底值

用法示例:
  python scripts/kpi_extract_output.py ^
    -i "门店KPI指标20260625-o.xlsx" ^
    -o "output_kpi.xlsx" ^
    -m 202606 ^
    -c 南宁
"""

import argparse
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import warnings

warnings.filterwarnings("ignore")


# ──────────────────────────────────────────────
# 大类名称别名映射
# 源数据"月度指标值" → 输出列"月度目标值"
# ──────────────────────────────────────────────
BIG_CATS_MAP = {
    "月度目标值": ["月度指标值", "月度目标值", "月度目标", "目标值"],
    "权重":       ["权重"],
    "封顶值":     ["封顶值", "封顶"],
    "保底值":     ["保底值", "保底"],
}


# ──────────────────────────────────────────────
# 数据解析
# ──────────────────────────────────────────────

def parse_data_sheet(filepath, data_sheet):
    """
    读取数据 sheet，解析双行表头，返回：
      - df_data：含有 (大类, 指标名) 多级列的 DataFrame（数据行）
      - info_cols：属于门店基础信息的列 key 列表（非指标列）
    """
    raw = pd.read_excel(filepath, sheet_name=data_sheet, header=None, dtype=str)

    if raw.shape[0] < 2:
        raise ValueError("Sheet '{}' 行数不足，无法解析双行表头。".format(data_sheet))

    # 第0行：大类（月度指标值/权重/封顶值...），空值前向填充
    big_cats = raw.iloc[0].ffill().tolist()
    # 第1行：列名（门店基础信息 或 指标名）
    col_names = raw.iloc[1].tolist()

    multi_cols = list(zip(big_cats, col_names))

    # 数据从第2行开始
    df_data = raw.iloc[2:].reset_index(drop=True)
    df_data.columns = pd.MultiIndex.from_tuples(multi_cols)

    # 找出基础信息列：第0行 NaN（大类为空）的列
    info_cols = []
    for big, col in multi_cols:
        if pd.isna(big) or str(big).strip() == "" or str(big).lower() == "nan":
            info_cols.append((big, col))

    return df_data, info_cols, multi_cols


def extract_unique_indicators(multi_cols, info_col_keys):
    """
    从多级列中提取唯一指标名列表。
    只保留在"月度指标值（或其别名）"大类下出现的指标，
    避免末尾被前向填充的非指标列（如类型合并、销售点）混入。
    """
    # 月度目标值的候选大类名
    target_big_candidates = BIG_CATS_MAP.get("月度目标值", ["月度指标值"])

    seen = set()
    result = []
    for big, ind in multi_cols:
        if (big, ind) in info_col_keys:
            continue
        # 只从"月度指标值"大类取指标名
        big_str = str(big).strip() if pd.notna(big) else ""
        if not any(c in big_str for c in target_big_candidates):
            continue
        name = str(ind).strip() if pd.notna(ind) else ""
        if name and name not in seen:
            result.append(name)
            seen.add(name)
    return result



def find_col_by_keyword(info_cols, keywords):
    """在基础信息列中按关键字模糊匹配，返回第一个命中的 (big, col) key。"""
    for kw in keywords:
        for big, col in info_cols:
            if col and kw in str(col):
                return (big, col)
    return None


# ──────────────────────────────────────────────
# 核心转换逻辑
# ──────────────────────────────────────────────

def _find_big_key(multi_cols, output_label, ind_name):
    """在多级列中查找指定输出大类下的指标列 key（通过别名映射）。"""
    candidates = BIG_CATS_MAP.get(output_label, [output_label])
    for big, col in multi_cols:
        if col == ind_name and any(c in str(big) for c in candidates):
            return (big, col)
    return None


def _safe_val(row, key):
    """安全取值，空/NaN 返回 0。"""
    if key is None:
        return 0
    try:
        v = row[key]
        if pd.isna(v) or str(v).strip() in ("", "nan", "None"):
            return 0
        return v
    except Exception:
        return 0


def build_output(df_data, info_cols, multi_cols, month, city, id_col_key, name_col_key):
    """将宽表 df_data 转为长表，每行 = 一个门店 × 一个指标。"""
    info_col_keys = set(info_cols)
    unique_indicators = extract_unique_indicators(multi_cols, info_col_keys)

    # 预先构建每个指标对应的大类 key（避免循环内重复查找）
    ind_keys = {}
    for ind_name in unique_indicators:
        ind_keys[ind_name] = {
            "月度目标值": _find_big_key(multi_cols, "月度目标值", ind_name),
            "权重":       _find_big_key(multi_cols, "权重",       ind_name),
            "封顶值":     _find_big_key(multi_cols, "封顶值",     ind_name),
            "保底值":     _find_big_key(multi_cols, "保底值",     ind_name),
        }

    records = []
    for _, row in df_data.iterrows():
        store_id   = str(row[id_col_key]).strip()   if id_col_key   else ""
        store_name = str(row[name_col_key]).strip() if name_col_key else ""

        # 跳过空行
        if not store_id or store_id in ("nan", "None", ""):
            continue

        for ind_name in unique_indicators:
            keys = ind_keys[ind_name]
            records.append({
                "账期月份":   month,
                "地市":       city,
                "店中商编码": store_id,
                "店中商名称": store_name,
                "指标名称":   ind_name,
                "月度目标值": _safe_val(row, keys["月度目标值"]),
                "权重":       _safe_val(row, keys["权重"]),
                "封顶值":     _safe_val(row, keys["封顶值"]),
                "保底值":     _safe_val(row, keys["保底值"]),
            })

    df_out = pd.DataFrame(records)

    # 数值列转换
    for col in ["月度目标值", "权重", "封顶值", "保底值"]:
        df_out[col] = pd.to_numeric(df_out[col], errors="coerce").fillna(0)

    return df_out


# ──────────────────────────────────────────────
# Excel 输出美化
# ──────────────────────────────────────────────

def write_styled_excel(df_out, output_path):
    """将 DataFrame 写出为带格式的 Excel 文件。"""
    df_out.to_excel(output_path, index=False, sheet_name="输出")

    wb = load_workbook(output_path)
    ws = wb.active

    header_fill = PatternFill("solid", start_color="366092", end_color="366092")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    data_font   = Font(name="Arial", size=10)
    center      = Alignment(horizontal="center", vertical="center")
    left        = Alignment(horizontal="left",   vertical="center")
    thin        = Side(style="thin", color="CCCCCC")
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    # 列宽：账期月份/地市/编码/名称/指标/目标值/权重/封顶/保底
    col_widths = [12, 8, 18, 22, 22, 12, 8, 10, 8]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    # 表头样式
    for cell in ws[1]:
        cell.fill      = header_fill
        cell.font      = header_font
        cell.alignment = center
        cell.border    = border

    num_cols = {6, 7, 8, 9}   # 月度目标值/权重/封顶值/保底值
    alt_fill = PatternFill("solid", start_color="EBF3F9", end_color="EBF3F9")

    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, cell in enumerate(row, 1):
            cell.font   = data_font
            cell.border = border
            if fill:
                cell.fill = fill
            if col_idx in num_cols:
                cell.alignment     = center
                cell.number_format = "General"
            else:
                cell.alignment = left

    ws.freeze_panes = "A2"
    wb.save(output_path)


# ──────────────────────────────────────────────
# 入口
# ──────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="从指标sheet宽表提取KPI数据，输出长表Excel"
    )
    parser.add_argument("-i", "--input",    required=True,  help="源 Excel 文件路径")
    parser.add_argument("-o", "--output",   required=True,  help="输出 Excel 文件路径")
    parser.add_argument("-m", "--month",    required=True,  help="账期月份，如 202606")
    parser.add_argument("-c", "--city",     required=True,  help="地市，如 南宁")
    parser.add_argument("--data_sheet",     default="指标", help="数据Sheet名（默认：指标）")
    parser.add_argument("--id_col",         default="店中商编码", help="编码列名（默认：店中商编码）")
    parser.add_argument("--name_col",       default="店中商名称", help="名称列名（默认：店中商名称）")
    args = parser.parse_args()

    # 1. 解析数据 sheet
    try:
        df_data, info_cols, multi_cols = parse_data_sheet(args.input, args.data_sheet)
    except Exception as e:
        print("[ERROR] 读取Sheet失败: {}".format(e), file=sys.stderr)
        sys.exit(1)

    # 2. 定位编码和名称列
    id_col_key   = find_col_by_keyword(info_cols, [args.id_col,   "编码", "商编码"])
    name_col_key = find_col_by_keyword(info_cols, [args.name_col, "名称", "店名"])

    if id_col_key is None:
        print("[WARN] 未找到编码列（关键字: {}），输出编码列将为空。".format(args.id_col))
    if name_col_key is None:
        print("[WARN] 未找到名称列（关键字: {}），输出名称列将为空。".format(args.name_col))

    # 3. 构建长表
    df_out = build_output(
        df_data, info_cols, multi_cols,
        month=str(args.month),
        city=str(args.city),
        id_col_key=id_col_key,
        name_col_key=name_col_key,
    )

    if df_out.empty:
        print("[WARN] 未提取到任何数据，请检查源文件结构。")
        sys.exit(0)

    # 4. 写出 Excel
    write_styled_excel(df_out, args.output)

    print("完成！共输出 {} 行数据 -> {}".format(len(df_out), args.output))
    print("账期月份: {}，地市: {}".format(args.month, args.city))
    print("涉及门店: {} 家，指标数: {} 个".format(
        df_out["店中商编码"].nunique(), df_out["指标名称"].nunique()
    ))


if __name__ == "__main__":
    main()
