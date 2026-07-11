"""
test_excel_xlookup.py
XLookup 合并脚本单元测试
"""
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook, Workbook

sys.path.insert(0, str(Path(__file__).parent.parent))
from scripts.excel_xlookup import xlookup_merge, get_columns


# ─── Fixtures ────────────────────────────────────────────────────

@pytest.fixture
def main_excel(tmp_path):
    """主文件：3 行数据，lookup 列为「编号」"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["编号", "姓名", "备注"])
    ws.append(["001", "张三", "备注A"])
    ws.append(["002", "李四", "备注B"])
    ws.append(["999", "王五", "备注C"])  # 999 在副文件中不存在
    fpath = tmp_path / "main.xlsx"
    wb.save(fpath)
    return str(fpath)


@pytest.fixture
def lookup_excel(tmp_path):
    """副文件：编号→部门、金额，其中 001 有两条记录"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["编号", "部门", "金额"])
    ws.append(["001", "财务部", 1000])
    ws.append(["001", "财务部", 2000])  # 重复 key：count=2，取第一条
    ws.append(["002", "技术部", 3000])
    fpath = tmp_path / "lookup.xlsx"
    wb.save(fpath)
    return str(fpath)


# ─── get_columns ──────────────────────────────────────────────────

class TestGetColumns:
    def test_basic(self, main_excel):
        cols = get_columns(main_excel)
        assert cols == ["编号", "姓名", "备注"]

    def test_empty_sheet_name_uses_active(self, main_excel):
        cols = get_columns(main_excel, sheet_name=None)
        assert cols == ["编号", "姓名", "备注"]


# ─── xlookup_merge ───────────────────────────────────────────────

class TestXlookupMerge:
    def test_basic_merge(self, main_excel, lookup_excel, tmp_path):
        """正常匹配：输出列 = 主文件列 + result_cols + 匹配记录数"""
        out = str(tmp_path / "out.xlsx")
        xlookup_merge(
            main_file=main_excel, main_sheet="",
            lookup_col="编号",
            lookup_file=lookup_excel, lookup_sheet="",
            match_col="编号",
            result_cols=["部门", "金额"],
            out_file=out,
        )
        wb = load_workbook(out)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert header == ["编号", "姓名", "备注", "部门", "金额", "匹配记录数"]
        wb.close()

    def test_matched_row_values(self, main_excel, lookup_excel, tmp_path):
        """匹配成功时：result_cols 值正确，记录数正确"""
        out = str(tmp_path / "out.xlsx")
        xlookup_merge(
            main_file=main_excel, main_sheet="",
            lookup_col="编号",
            lookup_file=lookup_excel, lookup_sheet="",
            match_col="编号",
            result_cols=["部门", "金额"],
            out_file=out,
        )
        wb = load_workbook(out)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        # 001 → 财务部, 1000, count=2
        assert rows[0][3] == "财务部"
        assert rows[0][4] == "1000"   # 强制文本
        assert rows[0][5] == "2"      # 重复 key 出现 2 次
        # 002 → 技术部, count=1
        assert rows[1][3] == "技术部"
        assert rows[1][5] == "1"
        wb.close()

    def test_unmatched_row_fills_empty(self, main_excel, lookup_excel, tmp_path):
        """未匹配行：result_cols 全为空，匹配记录数为 '0'"""
        out = str(tmp_path / "out.xlsx")
        xlookup_merge(
            main_file=main_excel, main_sheet="",
            lookup_col="编号",
            lookup_file=lookup_excel, lookup_sheet="",
            match_col="编号",
            result_cols=["部门", "金额"],
            out_file=out,
        )
        wb = load_workbook(out)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        # 第3行 999 未匹配
        assert rows[2][3] in ("", None)
        assert rows[2][4] in ("", None)
        assert rows[2][5] == "0"
        wb.close()

    def test_all_cells_text_format(self, main_excel, lookup_excel, tmp_path):
        """所有输出单元格 number_format 应为 '@'"""
        out = str(tmp_path / "out.xlsx")
        xlookup_merge(
            main_file=main_excel, main_sheet="",
            lookup_col="编号",
            lookup_file=lookup_excel, lookup_sheet="",
            match_col="编号",
            result_cols=["部门", "金额"],
            out_file=out,
        )
        wb = load_workbook(out)
        ws = wb.active
        for row in ws.iter_rows():
            for cell in row:
                assert cell.number_format == "@", f"{cell.coordinate}: {cell.number_format}"
        wb.close()

    def test_custom_count_col_name(self, main_excel, lookup_excel, tmp_path):
        """匹配记录数列名可自定义"""
        out = str(tmp_path / "out.xlsx")
        xlookup_merge(
            main_file=main_excel, main_sheet="",
            lookup_col="编号",
            lookup_file=lookup_excel, lookup_sheet="",
            match_col="编号",
            result_cols=["部门"],
            out_file=out,
            match_count_col_name="重复次数",
        )
        wb = load_workbook(out)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert "重复次数" in header
        wb.close()

    def test_row_count_preserved(self, main_excel, lookup_excel, tmp_path):
        """输出行数与主文件完全一致"""
        out = str(tmp_path / "out.xlsx")
        xlookup_merge(
            main_file=main_excel, main_sheet="",
            lookup_col="编号",
            lookup_file=lookup_excel, lookup_sheet="",
            match_col="编号",
            result_cols=["部门"],
            out_file=out,
        )
        wb = load_workbook(out)
        ws = wb.active
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(data_rows) == 3   # 主文件 3 行数据
        wb.close()

    def test_main_file_not_found(self, lookup_excel, tmp_path):
        with pytest.raises(SystemExit):
            xlookup_merge(
                main_file="/nonexistent/main.xlsx", main_sheet="",
                lookup_col="编号",
                lookup_file=lookup_excel, lookup_sheet="",
                match_col="编号", result_cols=["部门"],
                out_file=str(tmp_path / "out.xlsx"),
            )

    def test_empty_result_cols(self, main_excel, lookup_excel, tmp_path):
        with pytest.raises(SystemExit):
            xlookup_merge(
                main_file=main_excel, main_sheet="",
                lookup_col="编号",
                lookup_file=lookup_excel, lookup_sheet="",
                match_col="编号", result_cols=[],
                out_file=str(tmp_path / "out.xlsx"),
            )

    def test_lookup_col_not_in_main(self, main_excel, lookup_excel, tmp_path):
        """查找列不存在于主文件时应退出"""
        with pytest.raises(SystemExit):
            xlookup_merge(
                main_file=main_excel, main_sheet="",
                lookup_col="不存在的列",
                lookup_file=lookup_excel, lookup_sheet="",
                match_col="编号", result_cols=["部门"],
                out_file=str(tmp_path / "out.xlsx"),
            )

    def test_match_col_not_in_lookup(self, main_excel, lookup_excel, tmp_path):
        """对应列不存在于副文件时应抛出 ValueError"""
        with pytest.raises((SystemExit, ValueError)):
            xlookup_merge(
                main_file=main_excel, main_sheet="",
                lookup_col="编号",
                lookup_file=lookup_excel, lookup_sheet="",
                match_col="不存在的列", result_cols=["部门"],
                out_file=str(tmp_path / "out.xlsx"),
            )
