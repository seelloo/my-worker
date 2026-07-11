import pytest
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import os
import openpyxl
from pathlib import Path
from scripts.collect_excel_column import collect_column_data

def test_collect_column_data(tmp_path):
    src_dir = tmp_path / "src"
    dest_dir = tmp_path / "dest"
    src_dir.mkdir()
    dest_dir.mkdir()
    
    file_path = src_dir / "test_data.xlsx"
    wb = openpyxl.Workbook()
    
    # 构造第一个 Sheet: 有 P 列数据，包含表头和重复项
    ws1 = wb.active
    ws1.title = "Sheet1"
    # 我们故意放到 P (第16列)
    # openpyxl 行列从 1 开始
    ws1.cell(row=1, column=16, value="第一张表的表头应该被跳过") 
    ws1.cell(row=2, column=16, value="apple")
    ws1.cell(row=3, column=16, value="banana")
    ws1.cell(row=4, column=16, value="apple")      # 重复数据
    ws1.cell(row=5, column=16, value="")           # 空白数据
    ws1.cell(row=6, column=16, value=None)         # None数据
    
    # 构造第二个 Sheet: 也有 P 列数据，补充和交叉重复
    ws2 = wb.create_sheet(title="Sheet2")
    ws2.cell(row=1, column=16, value="第二张表的表头也应跳过")
    ws2.cell(row=2, column=16, value="orange")
    ws2.cell(row=3, column=16, value="banana")     # 与第一张表重复
    ws2.cell(row=4, column=16, value=" grape ")    # 带两边空格的数据，测试清理
    
    # 构造第三个 Sheet: 根本没有 P 列数据
    ws3 = wb.create_sheet(title="Sheet3")
    ws3.cell(row=1, column=1, value="这里只有A列数据，没有P列")
    
    wb.save(file_path)
    
    # 我们执行合并测试，提取P列
    out_file = dest_dir / "test_data.xlsx"
    target_sheet = "Final_Collection"
    
    collect_column_data(
        src_file=str(file_path),
        dest_path_str=str(dest_dir),
        target_column="P",
        out_sheet_name=target_sheet
    )
    
    # 验证有没有顺利另存出来
    assert out_file.exists()
    
    # 加载输出文件验证其逻辑
    wb_out = openpyxl.load_workbook(out_file)
    names = wb_out.sheetnames
    
    # 应该包含原本的三张表，以及最后那张新加的表
    assert target_sheet in names
    
    out_ws = wb_out[target_sheet]
    # 我们预期的去重非空数据应该有：
    # "apple", "banana", "orange", "grape" (四个)
    # 由于第一行被写死了 "汇总去重数据"
    assert out_ws.cell(row=1, column=1).value == "汇总去重数据"
    
    # 读取里面所有 A列 收集的值（从第二行开始）
    extracted_data = []
    for row in range(2, out_ws.max_row + 1):
        extracted_data.append(out_ws.cell(row=row, column=1).value)
        
    print("提取结果:", extracted_data)
        
    assert len(extracted_data) == 4
    assert set(extracted_data) == {"apple", "banana", "orange", "grape"}
    
    # 测试跳过表头
    assert "第一张表的表头应该被跳过" not in extracted_data
    assert "第二张表的表头也应跳过" not in extracted_data
    
    wb_out.close()
