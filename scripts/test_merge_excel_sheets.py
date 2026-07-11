import pytest
from pathlib import Path
import openpyxl
from utils.excel_utils import ExcelUtils
from scripts.merge_excel_sheets import merge_sheets

def test_get_valid_sheet_name():
    existing = set()
    name1 = ExcelUtils.get_valid_sheet_name("正常文件名", existing)
    assert name1 == "正常文件名"
    existing.add(name1)
    
    name2 = ExcelUtils.get_valid_sheet_name("文件:名称/测试\\1*?[].xlsx", existing)
    assert name2 == "文件_名称_测试_1____.xlsx"
    existing.add(name2)
    
    long_name = "这是一个非常非常非常非常非常非常非常非常非常长超出了三十一个字符长度限制的文件名"
    name3 = ExcelUtils.get_valid_sheet_name(long_name, existing)
    assert len(name3) <= 31
    existing.add(name3)
    
    name4 = ExcelUtils.get_valid_sheet_name(long_name, existing)
    assert len(name4) <= 31
    assert name4.endswith("_1")
    assert name4 != name3
    existing.add(name4)

def test_merge_sheets(tmp_path):
    src_dir = tmp_path / "src"
    src_dir.mkdir()
    out_file = tmp_path / "merged_output.xlsx"
    target_sheet = "DataSheet"
    
    # 辅助函数来创建测试的 Excel 文件
    def create_xlsx(filepath, sheet_name, data):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in data:
            ws.append(row)
        wb.save(filepath)

    data = [["ID", "Value"], [1, "A"], [2, "B"]]
    
    # 1. 正常具有该 Sheet 的 Excel
    file1 = src_dir / "FileOne.xlsx"
    create_xlsx(file1, target_sheet, data)
        
    # 2. 不具有该 Sheet 的 Excel (应该被跳过)
    file2 = src_dir / "FileTwo.xlsx"
    create_xlsx(file2, "WrongSheet", data)
        
    # 3. 位于子目录中，且有极长文件名，会引发截断
    sub_dir = src_dir / "sub"
    sub_dir.mkdir()
    long_name = "A"*40
    file3 = sub_dir / f"{long_name}.xlsx"
    create_xlsx(file3, target_sheet, data)
        
    # 4. 主目录下有同样长文件名的文件，会引发同名后缀逻辑
    file4 = src_dir / f"{long_name}.xlsx"
    create_xlsx(file4, target_sheet, data)
    
    # 执行合并
    merge_sheets(str(src_dir), str(out_file), target_sheet)
    
    assert out_file.exists(), "未生成输出文件"
    
    wb_out = openpyxl.load_workbook(out_file, data_only=True, read_only=True)
    names = wb_out.sheetnames
    
    # 期望合并了 3 个文件 (1, 3, 4)
    assert len(names) == 3
    
    assert "FileOne" in names
    
    # 因为 A*40 被截断到28个字符
    base_truncated = "A"*28
    
    # 至少有一个会叫 AAAAA... 或者加了后缀的
    assert any(n.startswith(base_truncated) for n in names)
    assert any(n.endswith("_1") for n in names)
    
    # 数据无损验证
    ws = wb_out["FileOne"]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3
    assert rows[2][1] == "B"
    wb_out.close()
