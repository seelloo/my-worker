"""
OFD Invoice Data Extractor (Web Integrated Version)
Based on direct XML value and object-reference (CustomTag) style OFD invoices parser.
"""

import zipfile, xml.etree.ElementTree as ET, re, os, sqlite3
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OFD_NS = 'http://www.ofdspec.org/2016'

for p, u in {'ofd': OFD_NS, 'invoice': 'http://www.chinatax.gov.cn/tirp'}.items():
    ET.register_namespace(p, u)

def read_zip(ofd_path, name):
    with zipfile.ZipFile(ofd_path, 'r') as z:
        return z.read(name).decode('utf-8', errors='replace')

def list_zip(ofd_path):
    with zipfile.ZipFile(ofd_path, 'r') as z:
        return z.namelist()

def parse_xml(text):
    try:
        return ET.fromstring(text)
    except ET.ParseError:
        clean = re.sub(r'\s+xmlns(?::\w+)?="[^"]+"', '', text)
        clean = re.sub(r'<(\w+):(\w+)', r'<\2', clean)
        clean = re.sub(r'</(\w+):(\w+)', r'</\2', clean)
        try: return ET.fromstring(clean)
        except ET.ParseError: return None

def local(tag):
    return tag.split('}')[-1] if '}' in tag else tag

def find_text(root, *tags):
    for elem in root.iter():
        if local(elem.tag) in tags:
            t = (elem.text or '').strip()
            if t: return t
    return ''

# ── Strategy 1: CustomTag object-reference style ────────────────────────────

def extract_via_customtag(ofd_path, names):
    """
    This OFD style stores field→ObjectRef mappings in CustomTag.xml,
    and actual text values as TextObject nodes in Page Content.xml.
    """
    ctag_file = next((n for n in names if n.endswith('CustomTag.xml')
                      and 'CustomTags' not in n), None)
    if not ctag_file:
        return None

    content_files = [n for n in names if re.search(r'Pages/Page_\d+/Content\.xml$', n)]
    if not content_files:
        return None

    id_text = {}
    for cf in content_files:
        try:
            root = parse_xml(read_zip(ofd_path, cf))
            if root is None: continue
            for obj in root.iter(f'{{{OFD_NS}}}TextObject'):
                oid = obj.get('ID')
                codes = [tc.text or '' for tc in obj.iter(f'{{{OFD_NS}}}TextCode')]
                text = ''.join(codes).strip()
                if oid and text:
                    id_text[oid] = text
        except Exception:
            pass

    if not id_text:
        return None

    custom_data = {}
    try:
        ofd_root = parse_xml(read_zip(ofd_path, 'OFD.xml'))
        if ofd_root is not None:
            for cd in ofd_root.iter(f'{{{OFD_NS}}}CustomData'):
                name_attr = cd.get('Name', '')
                val = (cd.text or '').strip()
                if name_attr and val:
                    custom_data[name_attr] = val
    except Exception:
        pass

    ctag_root = parse_xml(read_zip(ofd_path, ctag_file))
    if ctag_root is None:
        return None

    def refs_for(tag_local):
        vals = []
        for elem in ctag_root.iter():
            if local(elem.tag) == tag_local:
                for ref in elem.iter(f'{{{OFD_NS}}}ObjectRef'):
                    oid = (ref.text or '').strip()
                    if oid in id_text:
                        vals.append(id_text[oid])
        return vals

    def first(tag_local):
        vals = refs_for(tag_local)
        return vals[0] if vals else ''

    def joined(tag_local):
        return ''.join(refs_for(tag_local))

    items = []
    item_fields = ['Item','TaxScheme','MeasurementDimension','Amount','TaxAmount','Price','Quantity','Note']
    item_dict = {}
    for field in item_fields:
        vals = refs_for(field)
        if vals:
            item_dict[field] = vals[0]
    if item_dict:
        items.append(item_dict)

    note_val = first('Note')
    remarks = payee = reviewer = ''
    if note_val:
        m = re.search(r'收款人[：:]?\s*([^;；]+)', note_val)
        if m: payee = m.group(1).strip()
        m = re.search(r'复核人[：:]?\s*([^;；]+)', note_val)
        if m: reviewer = m.group(1).strip()
        remarks = note_val

    def clean_amount(tag):
        v = joined(tag).replace('¥','').replace(',','').strip()
        return v

    data = dict(
        file=os.path.basename(ofd_path),
        invoice_code   = custom_data.get('发票代码', ''),
        invoice_number = first('InvoiceNo') or custom_data.get('发票号码', ''),
        invoice_date   = first('IssueDate') or custom_data.get('开票日期', ''),
        invoice_type   = custom_data.get('发票类型', ''),
        check_code     = custom_data.get('校验码', ''),
        machine_number = custom_data.get('机器编号', ''),
        buyer_name     = first('BuyerName'),
        buyer_tax_id   = first('BuyerTaxID') or custom_data.get('购买方纳税人识别号', ''),
        buyer_address  = first('BuyerAddrTel'),
        buyer_bank     = first('BuyerBankAccount'),
        seller_name    = first('SellerName'),
        seller_tax_id  = first('SellerTaxID') or custom_data.get('销售方纳税人识别号', ''),
        seller_address = first('SellerAddrTel'),
        seller_bank    = first('SellerBankAccount'),
        amount_without_tax = clean_amount('TaxExclusiveTotalAmount') or custom_data.get('合计金额', ''),
        tax_amount         = clean_amount('TaxTotalAmount') or custom_data.get('合计税额', ''),
        total_amount       = clean_amount('TaxInclusiveTotalAmount'),
        total_in_words     = id_text.get('6934', '') or first('TotalAmountInWords'),
        remarks  = remarks,
        payee    = payee,
        reviewer = reviewer,
        drawer   = first('InvoiceClerk'),
        items    = items,
    )

    if not data['total_in_words']:
        mapped_ids = set()
        for elem in ctag_root.iter(f'{{{OFD_NS}}}ObjectRef'):
            mapped_ids.add((elem.text or '').strip())
        for oid, text in id_text.items():
            if oid not in mapped_ids and re.search(r'[圆元仟佰拾万零壹贰叁肆伍陆柒捌玖]', text):
                data['total_in_words'] = text
                break

    return data


# ── Strategy 2: Direct XML values ───────────────────────────────────────────

def extract_via_direct_xml(ofd_path, names):
    xml_files = [n for n in names if n.endswith('.xml')]
    best_root, best_score = None, 0
    for name in xml_files:
        root = parse_xml(read_zip(ofd_path, name))
        if root is None: continue
        score = sum(1 for _ in root.iter())
        if score > best_score:
            best_score, best_root = score, root
    if best_root is None:
        return None
    r = best_root

    def extract_items_direct(root):
        items = []
        for elem in root.iter():
            if local(elem.tag) in {'GoodInfo','GoodsInfo','Detail','InvoiceGoodsInfo'}:
                item = {local(c.tag): (c.text or '').strip() for c in elem}
                if item: items.append(item)
        return items

    return dict(
        file=os.path.basename(ofd_path),
        invoice_code   = find_text(r,'InvoiceCode','FPDaima','invoiceCode'),
        invoice_number = find_text(r,'InvoiceNo','FPHaoMa','invoiceNo','InvoiceNumber'),
        invoice_date   = find_text(r,'InvoiceDate','KaiPiaoRiQi','invoiceDate','IssueDate'),
        invoice_type   = find_text(r,'InvoiceType','FaPiaoLeiXing','invoiceType'),
        check_code     = find_text(r,'CheckCode','JiaoYanMa','checkCode'),
        machine_number = find_text(r,'MachineNo','JiQiBianHao','machineNo'),
        buyer_name     = find_text(r,'BuyerName','GouMaiFangMingCheng','buyerName'),
        buyer_tax_id   = find_text(r,'BuyerTaxID','GouMaiFangShuiHao','buyerTaxID','BuyerRegisterNum'),
        buyer_address  = find_text(r,'BuyerAddrTel','GouMaiFangDiZhiDianHua','buyerAddrTel'),
        buyer_bank     = find_text(r,'BuyerBankAccount','GouMaiFangKaiHuHangHaoMa','buyerBankAccount'),
        seller_name    = find_text(r,'SellerName','XiaoShouFangMingCheng','sellerName'),
        seller_tax_id  = find_text(r,'SellerTaxID','XiaoShouFangShuiHao','sellerTaxID','SellerRegisterNum'),
        seller_address = find_text(r,'SellerAddrTel','XiaoShouFangDiZhiDianHua','sellerAddrTel'),
        seller_bank    = find_text(r,'SellerBankAccount','XiaoShouFangKaiHuHangHaoMa','sellerBankAccount'),
        amount_without_tax = find_text(r,'TaxExclusiveAmount','HeJiJinE','taxExclusiveAmount','Amount'),
        tax_amount         = find_text(r,'TaxAmount','HeJiShuiE','taxAmount'),
        total_amount       = find_text(r,'TaxInclusiveAmount','JiaShuiHeJi','taxInclusiveAmount','TotalAmount'),
        total_in_words     = find_text(r,'TotalAmountInWords','JiaShuiHeJiDaXie','totalAmountInWords'),
        remarks  = find_text(r,'Remarks','BeiZhu','remarks','Note'),
        payee    = find_text(r,'Payee','ShouKuanRen','payee'),
        reviewer = find_text(r,'Reviewer','FuHe','reviewer'),
        drawer   = find_text(r,'Drawer','KaiPiaoRen','drawer'),
        items    = extract_items_direct(r),
    )


def extract_invoice_data(ofd_path):
    try:
        names = list_zip(ofd_path)
    except Exception as e:
        return dict(file=os.path.basename(ofd_path), _error=str(e))

    data = extract_via_customtag(ofd_path, names)
    if data and any(data.get(k) for k in ('invoice_number','buyer_name','seller_name')):
        return data

    data = extract_via_direct_xml(ofd_path, names)
    if data:
        return data

    return dict(file=os.path.basename(ofd_path), _error='Could not extract invoice data')


# ── SQL Database Persistence ───────────────────────────────────────────────

def init_db(db_path='invoices_archive.db'):
    conn = sqlite3.connect(db_path)
    c = conn.cursor()
    c.execute('''
        CREATE TABLE IF NOT EXISTS invoices (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_name TEXT,
            invoice_code TEXT,
            invoice_number TEXT,
            invoice_date TEXT,
            invoice_type TEXT,
            check_code TEXT,
            machine_number TEXT,
            buyer_name TEXT,
            buyer_tax_id TEXT,
            seller_name TEXT,
            seller_tax_id TEXT,
            amount_without_tax REAL,
            tax_amount REAL,
            total_amount REAL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(invoice_code, invoice_number)
        )
    ''')
    conn.commit()
    conn.close()

def save_to_sqlite(records, db_path='invoices_archive.db'):
    init_db(db_path)
    conn = sqlite3.connect(db_path)
    try:
        c = conn.cursor()
        
        insert_sql = '''
            INSERT OR IGNORE INTO invoices (
                file_name, invoice_code, invoice_number, invoice_date, invoice_type, 
                check_code, machine_number, buyer_name, buyer_tax_id, seller_name, 
                seller_tax_id, amount_without_tax, tax_amount, total_amount
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        '''
        
        inserted_count = 0
        for rec in records:
            if rec.get('_error'):
                continue
                
            def to_float(val):
                try: return float(str(val).replace(',', '').replace('¥', '').strip()) if val else 0.0
                except ValueError: return 0.0
                    
            c.execute(insert_sql, (
                rec.get('file', ''),
                rec.get('invoice_code', ''),
                rec.get('invoice_number', ''),
                rec.get('invoice_date', ''),
                rec.get('invoice_type', ''),
                rec.get('check_code', ''),
                rec.get('machine_number', ''),
                rec.get('buyer_name', ''),
                rec.get('buyer_tax_id', ''),
                rec.get('seller_name', ''),
                rec.get('seller_tax_id', ''),
                to_float(rec.get('amount_without_tax')),
                to_float(rec.get('tax_amount')),
                to_float(rec.get('total_amount'))
            ))
            
            # Cursor.rowcount tells us if the IGNORE bypassed the insertion
            if c.rowcount > 0:
                inserted_count += 1
                
        conn.commit()
    finally:
        conn.close()
    if inserted_count > 0:
        print(f"🗄️ 已静默沉淀 {inserted_count} 条新增发票记录至本地专属数据库 -> {db_path}")

# ── Excel output ─────────────────────────────────────────────────────────────

def _mk_fill(c): return PatternFill('solid', start_color=c)
def _mk_border(s='thin', c='C0C0C0'):
    sd = Side(style=s, color=c)
    return Border(left=sd, right=sd, top=sd, bottom=sd)

HDR_FILL   = _mk_fill('1F3864')
SEC_FILL   = _mk_fill('D6E4F0')
AMT_FILL   = _mk_fill('EBF5EB')
ALT_FILL   = _mk_fill('F5F9FF')
TTL_FILL   = _mk_fill('EBF0FA')
HDR_FONT   = Font(name='Arial', bold=True, color='FFFFFF', size=11)
SEC_FONT   = Font(name='Arial', bold=True, color='1F3864', size=10)
LBL_FONT   = Font(name='Arial', bold=True, color='404040', size=10)
VAL_FONT   = Font(name='Arial', color='000000', size=10)
TTL_FONT   = Font(name='Arial', bold=True, color='1F3864', size=13)
THIN       = _mk_border()
C_LEFT     = Alignment(horizontal='left', vertical='center', wrap_text=True)
C_CTR      = Alignment(horizontal='center', vertical='center')

def _c(ws, row, col, val='', font=None, fill=None, align=None, border=None, num_fmt=None):
    c = ws.cell(row=row, column=col, value=val)
    if font:    c.font = font
    if fill:    c.fill = fill
    if align:   c.alignment = align
    if border:  c.border = border
    if num_fmt: c.number_format = num_fmt
    return c

def _section(ws, row, label, ncols=2):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    _c(ws, row, 1, label, font=SEC_FONT, fill=SEC_FILL, align=C_LEFT, border=THIN)
    ws.row_dimensions[row].height = 18

def _field(ws, row, label, value, amount=False):
    _c(ws, row, 1, label, font=LBL_FONT, align=C_LEFT, border=THIN)
    fill = AMT_FILL if amount else None
    nfmt = '#,##0.00' if amount and isinstance(value, float) else None
    _c(ws, row, 2, value, font=VAL_FONT, fill=fill, align=C_LEFT, border=THIN, num_fmt=nfmt)
    ws.row_dimensions[row].height = 16

def save_xlsx(records, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Invoice Summary'
    ws.sheet_view.showGridLines = False

    headers = ['File','Invoice Code','Invoice No.','Date','Type',
               'Buyer Name','Buyer Tax ID','Seller Name','Seller Tax ID',
               'Amount (excl. tax)','Tax Amount','Total (incl. tax)','Drawer','Remarks']
    keys    = ['file','invoice_code','invoice_number','invoice_date','invoice_type',
               'buyer_name','buyer_tax_id','seller_name','seller_tax_id',
               'amount_without_tax','tax_amount','total_amount','drawer','remarks']
    amt_idx = {keys.index(k)+1 for k in ('amount_without_tax','tax_amount','total_amount')}

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    _c(ws,1,1,'OFD 发票提取汇总表', font=TTL_FONT, fill=TTL_FILL, align=C_CTR)
    ws.row_dimensions[1].height = 28

    for col, h in enumerate(headers, 1):
        _c(ws, 2, col, h, font=HDR_FONT, fill=HDR_FILL, align=C_CTR, border=THIN)
    ws.row_dimensions[2].height = 22

    for r, rec in enumerate(records, 3):
        fill = ALT_FILL if r % 2 == 0 else None
        for col, key in enumerate(keys, 1):
            val = rec.get(key, '')
            nfmt = None
            if col in amt_idx and val:
                try: val = float(val); nfmt = '#,##0.00'
                except ValueError: pass
            _c(ws, r, col, val, font=VAL_FONT, fill=fill, align=C_LEFT, border=THIN, num_fmt=nfmt)
        ws.row_dimensions[r].height = 16

    for col, w in enumerate([30,16,22,14,20,30,22,30,22,18,14,18,14,30], 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = 'A3'

    for rec in records:
        sname = re.sub(r'[\\/*?:\[\]]', '_', rec.get('invoice_number') or rec['file'])[:31]
        ws2 = wb.create_sheet(title=sname)
        ws2.sheet_view.showGridLines = False
        ws2.column_dimensions['A'].width = 26
        ws2.column_dimensions['B'].width = 44

        row = 1
        ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        _c(ws2, row, 1, f'发票 — {rec.get("invoice_number", rec["file"])}',
           font=TTL_FONT, fill=TTL_FILL, align=C_CTR)
        ws2.row_dimensions[row].height = 28; row += 1

        _section(ws2, row, '📄  发票基本信息'); row += 1
        for lbl, key in [('发票号码','invoice_number'),('开票日期','invoice_date'),
                          ('发票代码','invoice_code'),('发票类型','invoice_type'),
                          ('校验码','check_code'),('机器编号','machine_number')]:
            _field(ws2, row, lbl, rec.get(key,'')); row += 1

        row += 1
        _section(ws2, row, '🏢  购买方信息'); row += 1
        for lbl, key in [('购买方名称','buyer_name'),('纳税人识别号','buyer_tax_id'),
                          ('地址电话','buyer_address'),('开户行及账号','buyer_bank')]:
            _field(ws2, row, lbl, rec.get(key,'')); row += 1

        row += 1
        _section(ws2, row, '🏭  销售方信息'); row += 1
        for lbl, key in [('销售方名称','seller_name'),('纳税人识别号','seller_tax_id'),
                          ('地址电话','seller_address'),('开户行及账号','seller_bank')]:
            _field(ws2, row, lbl, rec.get(key,'')); row += 1

        row += 1
        _section(ws2, row, '💰  金额信息'); row += 1
        for lbl, key in [('合计金额（不含税）','amount_without_tax'),('合计税额','tax_amount'),
                          ('价税合计','total_amount'),('价税合计大写','total_in_words')]:
            val = rec.get(key,'')
            is_amt = key in ('amount_without_tax','tax_amount','total_amount')
            if is_amt and val:
                try: val = float(val)
                except ValueError: pass
            _field(ws2, row, lbl, val, amount=is_amt); row += 1

        row += 1
        _section(ws2, row, '📝  其他信息'); row += 1
        for lbl, key in [('备注','remarks'),('收款人','payee'),
                          ('复核人','reviewer'),('开票人','drawer')]:
            _field(ws2, row, lbl, rec.get(key,'')); row += 1

        items = rec.get('items', [])
        if items:
            row += 1
            all_keys = []
            for item in items:
                for k in item:
                    if k not in all_keys: all_keys.append(k)
            ncols = max(len(all_keys), 2)
            _section(ws2, row, f'🛒  商品明细 ({len(items)} 项)', ncols); row += 1
            for col, k in enumerate(all_keys, 1):
                ws2.column_dimensions[get_column_letter(col)].width = 18
                _c(ws2, row, col, k, font=LBL_FONT, fill=SEC_FILL, align=C_CTR, border=THIN)
            row += 1
            for item in items:
                for col, k in enumerate(all_keys, 1):
                    _c(ws2, row, col, item.get(k,''), font=VAL_FONT, align=C_LEFT, border=THIN)
                row += 1

        if rec.get('_error'):
            row += 1
            ws2.cell(row=row, column=1, value=f'⚠ 错误: {rec["_error"]}').font = Font(color='FF0000')

    wb.save(output_path)
    wb.close()
    print(f'✅ 已成功提取并保存发票数据至: {output_path}')


def process_ofd_extraction(src_dir: str, out_file: str):
    """
    性能加速版：并行化处理 OFD 文件抽提。
    """
    import concurrent.futures
    src_path = Path(src_dir).resolve()
    out_path = Path(out_file).resolve()
    
    if not src_path.exists() or not src_path.is_dir():
        print(f"❌ 源目录不存在: {src_dir}")
        return
        
    ofd_files = list(src_path.glob("*.ofd"))
    if not ofd_files:
        print(f"ℹ️ 在 {src_dir} 中未找到 .ofd 文件")
        return
        
    print(f"🚀 启动并发提取引擎 (共 {len(ofd_files)} 份 OFD)...")
    
    records = []
    # 使用线程池并发解析
    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as executor:
        future_to_ofd = {executor.submit(extract_invoice_data, str(f)): f for f in ofd_files}
        for future in concurrent.futures.as_completed(future_to_ofd):
            ofd_f = future_to_ofd[future]
            try:
                rec = future.result()
                records.append(rec)
                if rec.get('_error'):
                    print(f"  ⚠ {ofd_f.name} -> 提取抛错: {rec['_error']}")
                else:
                    print(f"  ✅ {ofd_f.name} -> 成功 ({rec.get('invoice_number','')})")
            except Exception as e:
                print(f"  ❌ {ofd_f.name} -> 线程崩溃: {e}")

    if not records:
        print("❌ 未提取到任何有效数据。")
        return
        
    out_path.parent.mkdir(parents=True, exist_ok=True)
    
    # Save to Local SQLite Database
    save_to_sqlite(records)
    
    # Generate the excel report
    save_xlsx(records, str(out_path))
if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="OFD 增值税发票全自动提取引擎")
    parser.add_argument("-s", "--src", required=True, help="OFD 票据源文件夹路径")
    parser.add_argument("-o", "--out", required=True, help="生成的汇总 Excel 文件路径")
    
    args = parser.parse_args()
    process_ofd_extraction(args.src, args.out)
