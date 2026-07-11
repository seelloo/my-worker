"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

test_excel_proportional_alloc.py
~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~~
单元测试：验证按比例分配脚本的核心逻辑。
"""

import os
import pytest
from openpyxl import Workbook, load_workbook

# 导入被测模块
from scripts.excel_proportional_alloc import excel_proportional_alloc


def _make_xlsx(rows: list, header: list, path: str):
    """辅助：创建测试用 Excel 文件"""
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()


def _read_xlsx(path: str) -> list:
    """辅助：读取结果 Excel，返回所有数据行（不含表头）"""
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    result = []
    for row in rows[1:]:  # 跳过 header
        result.append(row)
    return result


class TestProportionalAlloc:

    def test_basic_proportion(self, tmp_path):
        """核心用例：A=001 下 3 行权重 10/20/30，B=001 是 120，结果应为 20/40/60"""
        file_a = str(tmp_path / "a.xlsx")
        file_b = str(tmp_path / "b.xlsx")
        out = str(tmp_path / "out.xlsx")

        _make_xlsx(
            [["001", 10], ["001", 20], ["001", 30], ["002", 50]],
            ["编号", "金额"],
            file_a
        )
        _make_xlsx(
            [["001", 120], ["002", 100]],
            ["编号", "总额"],
            file_b
        )

        excel_proportional_alloc(
            file_a=file_a, file_b=file_b, output_path=out,
            key_a_name="编号", weight_a_name="金额",
            key_b_name="编号", value_b_name="总额",
            result_col="分配结果"
        )

        assert os.path.exists(out)
        data = _read_xlsx(out)
        assert len(data) == 4

        # 001 组：10/60*120=20, 20/60*120=40, 30/60*120=60
        assert data[0][2] == pytest.approx(20.0, abs=0.01)
        assert data[1][2] == pytest.approx(40.0, abs=0.01)
        assert data[2][2] == pytest.approx(60.0, abs=0.01)
        # 002 组：50/50*100 = 100
        assert data[3][2] == pytest.approx(100.0, abs=0.01)

    def test_unmatched_key_is_skipped(self, tmp_path):
        """A 中存在 B 不存在的 key，该行应该被直接跳过不输出"""
        file_a = str(tmp_path / "a.xlsx")
        file_b = str(tmp_path / "b.xlsx")
        out = str(tmp_path / "out.xlsx")

        _make_xlsx([["001", 10], ["999", 5]], ["编号", "金额"], file_a)
        _make_xlsx([["001", 100]], ["编号", "总额"], file_b)

        excel_proportional_alloc(
            file_a=file_a, file_b=file_b, output_path=out,
            key_a_name="编号", weight_a_name="金额",
            key_b_name="编号", value_b_name="总额",
            result_col="分配结果"
        )

        data = _read_xlsx(out)
        assert len(data) == 1
        assert data[0][2] == pytest.approx(100.0, abs=0.01)

    def test_custom_output_columns(self, tmp_path):
        """指定输出列：只输出 编号 和 分配结果"""
        file_a = str(tmp_path / "a.xlsx")
        file_b = str(tmp_path / "b.xlsx")
        out = str(tmp_path / "out.xlsx")

        _make_xlsx([["001", 10]], ["编号", "金额"], file_a)
        _make_xlsx([["001", 80]], ["编号", "总额"], file_b)

        excel_proportional_alloc(
            file_a=file_a, file_b=file_b, output_path=out,
            key_a_name="编号", weight_a_name="金额",
            key_b_name="编号", value_b_name="总额",
            result_col="分配结果",
            output_columns=["编号", "分配结果"]
        )

        data = _read_xlsx(out)
        assert len(data) == 1
        # 只有 2 列
        assert len([v for v in data[0] if v is not None]) == 2
        assert data[0][0] == "001"
        assert data[0][1] == pytest.approx(80.0, abs=0.01)

    def test_b_has_multiple_rows_same_key(self, tmp_path):
        """B 文件同一 key 多行：自动求和后分配"""
        file_a = str(tmp_path / "a.xlsx")
        file_b = str(tmp_path / "b.xlsx")
        out = str(tmp_path / "out.xlsx")

        _make_xlsx([["001", 1], ["001", 3]], ["编号", "金额"], file_a)
        _make_xlsx([["001", 60], ["001", 40]], ["编号", "总额"], file_b)

        excel_proportional_alloc(
            file_a=file_a, file_b=file_b, output_path=out,
            key_a_name="编号", weight_a_name="金额",
            key_b_name="编号", value_b_name="总额",
            result_col="分配结果"
        )

        data = _read_xlsx(out)
        # B 求和 = 100；A 权重 1:3，分别得 25 和 75
        assert data[0][2] == pytest.approx(25.0, abs=0.01)
        assert data[1][2] == pytest.approx(75.0, abs=0.01)

    def test_decimal_places(self, tmp_path):
        """验证小数位数参数生效"""
        file_a = str(tmp_path / "a.xlsx")
        file_b = str(tmp_path / "b.xlsx")
        out = str(tmp_path / "out.xlsx")

        _make_xlsx([["001", 1], ["001", 2]], ["编号", "金额"], file_a)
        _make_xlsx([["001", 10]], ["编号", "总额"], file_b)

        excel_proportional_alloc(
            file_a=file_a, file_b=file_b, output_path=out,
            key_a_name="编号", weight_a_name="金额",
            key_b_name="编号", value_b_name="总额",
            result_col="分配结果",
            decimal_places=0
        )

        data = _read_xlsx(out)
        # 1/3*10 ~ 3.33 -> round to 0 -> 3.0
        assert data[0][2] == pytest.approx(3.0, abs=0.5)

    def test_rounding_correction_sum_equals_target(self, tmp_path):
        """
        尾差修正核心验证：
        权重 1:1:1，B 值为 10 → 每份精确值 3.333...
        保留 2 位后各行 3.33，合计 9.99 ≠ 10
        修正后组内合计必须严格等于 10.00
        """
        file_a = str(tmp_path / "a.xlsx")
        file_b = str(tmp_path / "b.xlsx")
        out = str(tmp_path / "out.xlsx")

        _make_xlsx(
            [["001", 1], ["001", 1], ["001", 1]],
            ["编号", "金额"],
            file_a
        )
        _make_xlsx([["001", 10]], ["编号", "总额"], file_b)

        excel_proportional_alloc(
            file_a=file_a, file_b=file_b, output_path=out,
            key_a_name="编号", weight_a_name="金额",
            key_b_name="编号", value_b_name="总额",
            result_col="分配结果",
            decimal_places=2
        )

        data = _read_xlsx(out)
        assert len(data) == 3
        total = sum(row[2] for row in data)
        # 修正后合计必须严格等于原始 B 值 10
        assert total == pytest.approx(10.0, abs=1e-9)

    def test_rounding_correction_multiple_groups(self, tmp_path):
        """
        多组尾差修正：每组独立修正，每组合计必须等于对应 B 值
        """
        file_a = str(tmp_path / "a.xlsx")
        file_b = str(tmp_path / "b.xlsx")
        out = str(tmp_path / "out.xlsx")

        # 001: 3行权重 1:1:1, B=10 → 每组sum必须=10
        # 002: 2行权重 1:2, B=7  → 每组sum必须=7
        _make_xlsx(
            [["001", 1], ["001", 1], ["001", 1],
             ["002", 1], ["002", 2]],
            ["编号", "金额"],
            file_a
        )
        _make_xlsx([["001", 10], ["002", 7]], ["编号", "总额"], file_b)

        excel_proportional_alloc(
            file_a=file_a, file_b=file_b, output_path=out,
            key_a_name="编号", weight_a_name="金额",
            key_b_name="编号", value_b_name="总额",
            result_col="分配结果",
            decimal_places=2
        )

        data = _read_xlsx(out)
        assert len(data) == 5
        # 001 组合计 = 10
        sum_001 = sum(row[2] for row in data[:3])
        assert sum_001 == pytest.approx(10.0, abs=1e-9)
        # 002 组合计 = 7
        sum_002 = sum(row[2] for row in data[3:])
        assert sum_002 == pytest.approx(7.0, abs=1e-9)

