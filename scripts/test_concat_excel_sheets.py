import pytest
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pathlib import Path
from openpyxl import Workbook, load_workbook
from scripts.concat_excel_sheets import get_unique_path, concat_excel_folder

def test_get_unique_path(tmp_path):
    target = tmp_path / "test.xlsx"
    assert get_unique_path(target) == target
    target.touch()
    assert get_unique_path(target) == tmp_path / "test_1.xlsx"
    (tmp_path / "test_1.xlsx").touch()
    assert get_unique_path(target) == tmp_path / "test_2.xlsx"

def setup_test_excel(file_path: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["H1", "H2", "H3"])
    ws1.append(["A", "B", "C"])
    ws1.append(["D", "E", "F"])
    ws2 = wb.create_sheet(title="Sheet2")
    ws2.append(["H1", "H2", "H3"])
    ws2.append(["X", "Y", "Z"])
    ws2.append([None, "", None]) # 空行
    ws2.append(["M", "N", "O"])
    wb.create_sheet(title="Sheet3")
    wb.save(file_path)

def test_concat_excel_folder(tmp_path):
    src_dir = tmp_path / "src"
    src_dir.mkdir()
    setup_test_excel(src_dir / "a.xlsx")
    out_file = tmp_path / "out.xlsx"
    concat_excel_folder(str(src_dir), str(out_file), start_row=2)
    assert out_file.exists()
