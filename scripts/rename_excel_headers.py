import argparse
import os
import openpyxl

def load_mapping(mapping_file, old_col_idx=0, new_col_idx=1):
    """
    加载映射关系。
    默认取前两列：第一列为旧名(XXX)，第二列为新名(AAA)。
    """
    mapping = {}
    try:
        wb = openpyxl.load_workbook(mapping_file, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(min_row=1, values_only=True):
            if len(row) > max(old_col_idx, new_col_idx):
                old_val = str(row[old_col_idx]).strip() if row[old_col_idx] is not None else ""
                new_val = str(row[new_col_idx]).strip() if row[new_col_idx] is not None else ""
                if old_val:
                    mapping[old_val] = new_val
        wb.close()
        return mapping
    except Exception as e:
        print(f"Error loading mapping file: {e}")
        return {}

def rename_headers(src_path, mapping, out_path):
    """
    重命名 Excel 文件中的表头。
    """
    wb = None
    try:
        wb = openpyxl.load_workbook(src_path)
        modified = False
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            # 获取第一行内容
            first_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=False))
            if not first_row:
                continue
            
            headers = first_row[0]
            for cell in headers:
                val = str(cell.value).strip() if cell.value is not None else ""
                if val in mapping:
                    cell.value = mapping[val]
                    modified = True
                    
        if modified:
            os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
            wb.save(out_path)
            print(f"Successfully processed: {src_path} -> {out_path}")
        else:
            print(f"No matching headers found in: {src_path}")
            os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
            wb.save(out_path)
            
    except Exception as e:
        print(f"Error processing {src_path}: {e}")
    finally:
        if wb:
            wb.close()

def main():
    parser = argparse.ArgumentParser(description="Rename Excel headers based on a mapping file.")
    parser.add_argument("-s", "--src", required=True, help="Source Excel file or directory.")
    parser.add_argument("-m", "--mapping", required=True, help="Mapping Excel file path.")
    parser.add_argument("-o", "--out", required=True, help="Output file or directory path.")
    parser.add_argument("--old_idx", type=int, default=0, help="0-based index of old column in mapping file.")
    parser.add_argument("--new_idx", type=int, default=1, help="0-based index of new column in mapping file.")
    
    args = parser.parse_args()
    
    mapping = load_mapping(args.mapping, args.old_idx, args.new_idx)
    if not mapping:
        print("Empty or invalid mapping. Exiting.")
        return

    print(f"Loaded {len(mapping)} mapping rules.")

    if os.path.isdir(args.src):
        # Process directory
        os.makedirs(args.out, exist_ok=True)
        for filename in os.listdir(args.src):
            if filename.endswith(".xlsx") or filename.endswith(".xlsm"):
                src_file = os.path.join(args.src, filename)
                out_file = os.path.join(args.out, filename)
                rename_headers(src_file, mapping, out_file)
    else:
        # Process single file
        rename_headers(args.src, mapping, args.out)

if __name__ == "__main__":
    main()
