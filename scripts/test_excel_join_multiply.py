import pytest
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pathlib import Path
from openpyxl import Workbook, load_workbook
from scripts.excel_join_multiply import excel_join_multiply

def create_excel(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)
    wb.close()

def test_excel_join_multiply(tmp_path):
    # 1. 构造测试数据
    file_a = tmp_path / "file_a.xlsx"
    file_b = tmp_path / "file_b.xlsx"
    file_c = tmp_path / "file_c.xlsx"
    
    headers_a = ['id', 'a', 'c']
    rows_a = [
        [1, 'A1', 10],
        [2, 'A2', 20],
        [3, 'A3', 30]
    ]
    
    headers_b = ['b', 'c']
    rows_b = [
        ['A1', 2],
        ['A2', 3],
        ['A4', 4]
    ]
    
    create_excel(file_a, headers_a, rows_a)
    create_excel(file_b, headers_b, rows_b)
    
    # 2. 执行函数
    excel_join_multiply(str(file_a), str(file_b), str(file_c), 
                        key_a_name='a', key_b_name='b', col_a_name='c', col_b_name='c')
    
    # 3. 验证结果
    assert file_c.exists()
    wb_res = load_workbook(file_c, data_only=True)
    ws_res = wb_res.active
    
    # 获取所有行
    data = list(ws_res.iter_rows(values_only=True))
    header = data[0]
    rows = data[1:]
    
    assert header == ('id', 'a', 'c', '乘积结果')
    assert len(rows) == 3
    
    # A1: 10 * 2 = 20
    # A2: 20 * 3 = 60
    # A3: 30 * None = None (NaN in Excel)
    
    row_1 = next(r for r in rows if r[1] == 'A1')
    row_2 = next(r for r in rows if r[1] == 'A2')
    row_3 = next(r for r in rows if r[1] == 'A3')
    
    assert row_1[3] == 20
    assert row_2[3] == 60
    assert row_3[3] is None
    
    wb_res.close()

def test_excel_join_multiply_different_cols(tmp_path):
    # 1. 构造测试数据
    file_a = tmp_path / "file_a_diff.xlsx"
    file_b = tmp_path / "file_b_diff.xlsx"
    file_c = tmp_path / "file_c_diff.xlsx"
    
    headers_a = ['key_a', 'val_a']
    rows_a = [
        ['X', 100],
        ['Y', 200]
    ]
    
    headers_b = ['key_b', 'val_b']
    rows_b = [
        ['X', 5],
        ['Z', 10]
    ]
    
    create_excel(file_a, headers_a, rows_a)
    create_excel(file_b, headers_b, rows_b)
    
    # 2. 执行函数
    excel_join_multiply(str(file_a), str(file_b), str(file_c), 
                        key_a_name='key_a', key_b_name='key_b', 
                        col_a_name='val_a', col_b_name='val_b',
                        result_col='Result')
    
    # 3. 验证结果
    wb_res = load_workbook(file_c, data_only=True)
    ws_res = wb_res.active
    data = list(ws_res.iter_rows(values_only=True))
    
    header = data[0]
    assert 'Result' in header
    idx_res = header.index('Result')
    
    rows = data[1:]
    row_x = next(r for r in rows if r[0] == 'X')
    assert row_x[idx_res] == 500
    
    wb_res.close()
