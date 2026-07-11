"""
Excel处理工具类
提供统一的Excel文件操作功能，包括Sheet名称处理、列数据收集等
"""

import re
from pathlib import Path
from typing import List, Set, Optional, Union
import openpyxl
import xlrd


class ExcelUtils:
    """Excel处理工具类"""

    @staticmethod
    def get_valid_sheet_name(base_name: str, existing_names: Set[str]) -> str:
        """清理和缩短名称，确保它是一个有效的 Excel Sheet 名称，并且不冲突"""
        invalid_chars = r'[\\/*?:\[\]]'
        clean_name = re.sub(invalid_chars, '_', base_name)
        if not clean_name.strip():
            clean_name = "Sheet"

        max_len = 28
        current_name = clean_name[:max_len].strip()

        if current_name not in existing_names and len(current_name) <= 31:
            return current_name

        counter = 1
        while True:
            suffix = f"_{counter}"
            allowed_base_len = 31 - len(suffix)
            candidate = f"{clean_name[:allowed_base_len].strip()}{suffix}"

            if candidate not in existing_names:
                return candidate
            counter += 1

    @staticmethod
    def generate_new_sheet_name(original_name: str, remove_str: str, add_str: str, add_position: str, existing_names: Set[str]) -> str:
        """
        根据规则生成新的 Sheet 名称。
        优先处理去除，然后再处理增加。
        强行从后往前截断至 31 字，并解决同名冲突。
        """
        new_name = original_name

        # 1. 优先执行去除
        if remove_str:
            new_name = new_name.replace(remove_str, "")

        # 如果去除后为空，给一个默认底名
        if not new_name.strip():
            new_name = "Sheet"

        # 2. 其次执行增加
        if add_str:
            if add_position == "front":
                new_name = f"{add_str}{new_name}"
            elif add_position == "back":
                new_name = f"{new_name}{add_str}"

        # 3. 处理 31 个字符的硬性限制截断
        needs_warning = False

        # 我们保留最大28位长度，留3位给 _1 ~ _99，保证一定不会超过 31
        max_safe_len = 28
        if len(new_name) > 31:
            new_name = new_name[:max_safe_len]
            needs_warning = True

        # 4. 解决同名冲突
        candidate_name = new_name[:31]  # 如果没冲突，直接可以用最长31位的
        if candidate_name not in existing_names:
            if needs_warning:
                print(f"      ⚠️ 因新名字({original_name} -> ...)超过31个字，已从后往前截断为: '{candidate_name}'")
            return candidate_name

        # 如果有冲突，则使用截断到28位的基础上逐渐追加后缀
        counter = 1
        base_name_for_conflict = new_name[:max_safe_len]
        while True:
            suffix = f"_{counter}"
            # 再次确保后缀加上后总体 <= 31
            allowed_base_len = 31 - len(suffix)
            candidate = f"{base_name_for_conflict[:allowed_base_len]}{suffix}"

            if candidate not in existing_names:
                if needs_warning:
                    print(f"      ⚠️ 因新名字({original_name} -> ...)超过31个字，并发生重名，已处理为: '{candidate}'")
                return candidate
            counter += 1

    @staticmethod
    def process_xlsx_gen(file_path: Path, target_sheet: str):
        """
        生成器版：提取 xlsx 文件数据内容
        使用 read_only 模式，实时产发行数据。
        """
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        if target_sheet not in wb.sheetnames:
            wb.close()
            return None

        ws = wb[target_sheet]
        def row_generator():
            for row in ws.iter_rows(values_only=True):
                yield list(row)
            wb.close()

        return row_generator()

    @staticmethod
    def process_xls(file_path: Path, target_sheet: str):
        """
        xls 格式（老格式）提取，暂不做生成器优化（xlrd 限制）。
        """
        wb = xlrd.open_workbook(str(file_path))
        if target_sheet not in wb.sheet_names():
            return None
        ws = wb.sheet_by_name(target_sheet)
        rows = []
        for rx in range(ws.nrows):
            rows.append(ws.row_values(rx))
        return rows

    @staticmethod
    def collect_column_data_from_sheet(ws, target_column: str) -> Set[str]:
        """
        从单个Sheet中收集指定列的数据
        返回去重后的数据集合
        """
        collected_data = set()
        total_cells_scanned = 0

        try:
            col_cells = ws[target_column]
        except Exception:
            # 说明这列超出了现存最大列范围之类的情况
            return collected_data

        # 如果是单一单元格对象（比如只有1行数据），把它变成列表结构统一处理
        if not isinstance(col_cells, tuple):
            col_cells = (col_cells,)

        # 强制跳过第一行 (index=0)，也就是跳过表头
        if len(col_cells) > 1:
            for cell in col_cells[1:]:
                val = cell.value
                total_cells_scanned += 1

                # 我们只收集非空的内容（None 或者空白字符串过滤）
                if val is not None and str(val).strip() != "":
                    collected_data.add(str(val).strip())

        return collected_data

    @staticmethod
    def create_output_workbook() -> openpyxl.Workbook:
        """创建输出工作簿，使用write_only模式减少内存占用"""
        return openpyxl.Workbook(write_only=True)

    @staticmethod
    def create_output_sheet(workbook: openpyxl.Workbook, sheet_name: str) -> openpyxl.worksheet.worksheet.Worksheet:
        """在输出工作簿中创建新Sheet"""
        return workbook.create_sheet(title=sheet_name)

    @staticmethod
    def save_workbook(workbook: openpyxl.Workbook, output_path: Path) -> None:
        """保存工作簿到指定路径"""
        try:
            workbook.save(output_path)
            workbook.close()
        except Exception as e:
            raise Exception(f"无法保存文件 '{output_path.name}': {e}")